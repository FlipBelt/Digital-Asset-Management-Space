"""Account containment, employee binding and evidence-based offboarding on disposable PG."""

from datetime import UTC, date, datetime, timedelta
from io import BytesIO
from unittest.mock import MagicMock
from uuid import UUID, uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.config import get_settings
from app.db.session import SessionLocal
from app.main import app
from app.models import (
    AccessGrant,
    Account,
    Asset,
    AssetPlatformLink,
    AssetResponsibility,
    AssetType,
    AuditLog,
    Department,
    DingTalkPersonProfile,
    LegalEntity,
    Person,
    PlatformTenant,
    RegistrationIdentityProfile,
    ServiceProduct,
    User,
    UserRoleScope,
    UserSession,
)
from app.services.dingtalk import DingTalkDirectorySync, SyncResult
from tests.test_asset_center_integration import membership
from tests.test_subscription_verification import company_account as company_account

pytest_plugins = ["tests.test_asset_center_integration"]


def child(actors, company, **changes):
    body = {
        "request_id": str(uuid4()),
        "display_name": "隔离员工席位",
        "login_identifier": "seat-" + uuid4().hex + "@example.test",
        **changes,
    }
    response = actors.admin.post(
        f"/api/v1/workspace/platform-accounts/{company['asset']}/child-accounts",
        json=body,
    )
    assert response.status_code == 201, response.text
    return response.json(), body


def binding_body(actors, row, **changes):
    asset = actors.admin.get(f"/api/v1/assets/{row['asset_id']}").json()
    return {
        "request_id": str(uuid4()),
        "version": asset["version"],
        "primary_person_id": str(actors.person["owner"]),
        **changes,
    }


def binding(client, company, row, body):
    return client.patch(
        f"/api/v1/workspace/platform-accounts/{company['asset']}/child-accounts/{row['id']}/employee",
        json=body,
    )


def departure_body(actors, person="owner", **changes):
    with SessionLocal() as db:
        row = db.get(Person, actors.person[person])
        timestamp = row.updated_at.isoformat()
    return {
        "request_id": str(uuid4()),
        "expected_updated_at": timestamp,
        "employment_status": "departed",
        "departed_on": date.today().isoformat(),
        "evidence_note": "隔离测试：管理员已核对真实离职通知。",
        **changes,
    }


def departure(client, actors, body, person="owner"):
    return client.post(f"/api/v1/people/{actors.person[person]}/lifecycle", json=body)


def test_library_filters_before_count_pages_export_and_preserves_uuid(actors, company_account):
    row, _ = child(actors, company_account)
    with SessionLocal() as db:
        asset = db.get(Asset, UUID(company_account["asset"]))
        marker = "Library-" + uuid4().hex
        asset.name = marker + " parent"
        normal_type = db.scalar(select(AssetType.id).where(AssetType.code == "cloud_server"))
        # Profiles also exclude old accounts even if their category was mistyped.
        db.get(Asset, UUID(row["asset_id"])).asset_type_id = normal_type
        db.get(Asset, UUID(row["asset_id"])).name = marker + " child"
        regular = Asset(
            name=marker + " resource",
            asset_code="LIB-" + uuid4().hex,
            asset_type_id=normal_type,
            legal_entity_id=UUID(actors.entity),
        )
        identity = Asset(
            name=marker + " identity",
            asset_code="ID-" + uuid4().hex,
            asset_type_id=db.scalar(
                select(AssetType.id).where(AssetType.code == "registration_identity")
            ),
        )
        db.add_all([regular, identity])
        db.commit()
        regular_id = str(regular.id)
    result = actors.admin.get(
        "/api/v1/assets", params={"keyword": marker, "library_only": True, "page_size": 1}
    ).json()
    assert result["pagination"]["total"] == 1
    assert [item["id"] for item in result["data"]] == [regular_id]
    assert actors.admin.get(f"/api/v1/assets/{row['asset_id']}").status_code == 200
    assert (
        actors.admin.get("/api/v1/assets", params={"keyword": marker}).json()["pagination"]["total"]
        == 4
    )
    export = actors.admin.get("/api/v1/exports/assets.xlsx?library_only=true")
    names = {
        cell[1]
        for cell in load_workbook(BytesIO(export.content)).active.iter_rows(
            min_row=2, values_only=True
        )
    }
    assert marker + " resource" in names
    assert marker + " parent" not in names and marker + " child" not in names


def test_effective_company_owner_and_child_followup_are_one_parent(actors, company_account):
    row, _ = child(actors, company_account)
    with SessionLocal() as db:
        db.add(
            AssetResponsibility(
                asset_id=UUID(company_account["asset"]),
                person_id=actors.person["owner"],
                role_type="responsible",
                is_primary=True,
            )
        )
        db.commit()
    items = actors.admin.get("/api/v1/hudu/expirations").json()["items"]
    parent = next(item for item in items if item["id"] == company_account["asset"])
    assert parent["has_owner"] and parent["responsible_person_name"] == "测试身份 owner"
    assert parent["pending_child_count"] == 1
    assert row["asset_id"] not in {item["id"] for item in items}
    assert binding(actors.admin, company_account, row, binding_body(actors, row)).status_code == 200
    items = actors.admin.get("/api/v1/hudu/expirations").json()["items"]
    assert company_account["asset"] not in {item["id"] for item in items}


@pytest.mark.parametrize(
    "condition", ["department_only", "expired", "future", "departed", "proposed"]
)
def test_followup_does_not_count_invalid_owners(actors, company_account, condition):
    with SessionLocal() as db:
        asset = db.get(Asset, UUID(company_account["asset"]))
        asset.owner_department_id = actors.teams[0]
        if condition != "department_only":
            db.add(
                AssetResponsibility(
                    asset_id=asset.id,
                    person_id=actors.person["owner"],
                    role_type="proposed_responsible" if condition == "proposed" else "responsible",
                    is_primary=True,
                    ends_at=date.today() - timedelta(days=1) if condition == "expired" else None,
                    starts_at=date.today() + timedelta(days=1) if condition == "future" else None,
                )
            )
        if condition == "departed":
            db.get(Person, actors.person["owner"]).employment_status = "departed"
        db.commit()
    item = next(
        row
        for row in actors.admin.get("/api/v1/hudu/expirations").json()["items"]
        if row["id"] == company_account["asset"]
    )
    assert not item["has_owner"] and item["responsibility_issue"]


def test_hierarchy_contains_children_and_respects_personal_visibility(actors, company_account):
    row, _ = child(actors, company_account, primary_person_id=str(actors.person["owner"]))
    asset, _ = membership(actors.owner, actors)
    with SessionLocal() as db:
        db.get(ServiceProduct, UUID(actors.product)).platform_id = company_account["platform"]
        identity = Asset(
            name="不可猜测的平台身份",
            asset_code="IDENT-" + uuid4().hex,
            asset_type_id=db.scalar(
                select(AssetType.id).where(AssetType.code == "registration_identity")
            ),
        )
        db.add(identity)
        db.flush()
        db.add(
            RegistrationIdentityProfile(
                asset_id=identity.id,
                identity_type="email",
                identifier_value="identity@example.test",
                identifier_masked="i***@example.test",
                identifier_fingerprint=uuid4().hex,
                source_nature="unknown",
                verification_status="pending",
            )
        )
        db.commit()
    owner = actors.owner.get("/api/v1/account-hierarchy").json()
    platform = next(
        item for item in owner["platforms"] if item["id"] == str(company_account["platform"])
    )
    assert platform["company_accounts"][0]["children"][0]["id"] == row["id"]
    assert platform["company_accounts"][0]["children"][0]["person_name"] == "测试身份 owner"
    assert asset["id"] in {item["asset_id"] for item in platform["personal_accounts"]}
    other = actors.other.get("/api/v1/account-hierarchy").json()
    assert asset["id"] not in {
        item["asset_id"] for plat in other["platforms"] for item in plat["personal_accounts"]
    }
    assert any(item["name"] == "i***@example.test" for item in owner["unlinked_identities"])


def test_child_create_employee_binding_retry_and_no_new_grants(actors, company_account):
    row, body = child(actors, company_account, primary_person_id=str(actors.person["owner"]))
    retry = actors.admin.post(
        f"/api/v1/workspace/platform-accounts/{company_account['asset']}/child-accounts", json=body
    )
    assert retry.status_code == 201 and retry.json()["id"] == row["id"]
    changed = actors.admin.post(
        f"/api/v1/workspace/platform-accounts/{company_account['asset']}/child-accounts",
        json={**body, "display_name": "改变内容"},
    )
    assert changed.status_code == 409
    with SessionLocal() as db:
        assert not db.scalar(
            select(AccessGrant.id).where(AccessGrant.account_id == UUID(row["id"]))
        )
        assert (
            db.scalar(
                select(func.count())
                .select_from(AuditLog)
                .where(
                    AuditLog.action == "platform_account.child_account.create",
                    AuditLog.request_id == body["request_id"],
                )
            )
            == 1
        )


@pytest.mark.parametrize(
    "invalid", ["inactive", "departed", "foreign_company", "stale", "foreign_parent", "employee"]
)
def test_binding_rejects_invalid_employee_versions_and_permissions(
    actors, company_account, invalid
):
    row, _ = child(actors, company_account)
    body = binding_body(actors, row)
    with SessionLocal() as db:
        person = db.get(Person, actors.person["owner"])
        if invalid in {"inactive", "departed"}:
            person.employment_status = invalid
        elif invalid == "foreign_company":
            entity = LegalEntity(code="OTHER-" + uuid4().hex, name="其他隔离公司")
            db.add(entity)
            db.flush()
            person.legal_entity_id = entity.id
        db.commit()
    if invalid == "stale":
        body["version"] += 2
    if invalid == "foreign_parent":
        company_account = {**company_account, "asset": str(uuid4())}
    response = binding(
        actors.other if invalid == "employee" else actors.admin, company_account, row, body
    )
    expected = (
        403
        if invalid == "employee"
        else 404
        if invalid == "foreign_parent"
        else 409
        if invalid == "stale"
        else 422
    )
    assert response.status_code == expected, response.text


def test_existing_binding_updates_once_and_requires_csrf(actors, company_account):
    row, _ = child(actors, company_account)
    body = binding_body(actors, row)
    result = binding(actors.admin, company_account, row, body)
    assert (
        result.status_code == 200
        and result.json()["primary_person_id"] == body["primary_person_id"]
    )
    assert binding(actors.admin, company_account, row, body).json() == result.json()
    assert (
        binding(actors.admin, company_account, row, {**body, "primary_person_id": None}).status_code
        == 409
    )
    browser = TestClient(app)
    try:
        browser.cookies.set(
            get_settings().session_cookie_name, actors.admin.headers["X-PM-Session"]
        )
        assert binding(browser, company_account, row, body).status_code == 403
    finally:
        browser.close()


def test_departure_disables_only_target_login_and_preserves_handover(actors, company_account):
    row, _ = child(actors, company_account, primary_person_id=str(actors.person["owner"]))
    with SessionLocal() as db:
        db.add(
            AssetResponsibility(
                asset_id=UUID(company_account["asset"]),
                person_id=actors.person["owner"],
                role_type="responsible",
                is_primary=True,
            )
        )
        grant = AccessGrant(
            account_id=UUID(row["id"]),
            person_id=actors.person["owner"],
            grant_type="direct",
            status="active",
        )
        db.add(grant)
        db.commit()
        grant_id = grant.id
        user = db.scalar(select(User).where(User.person_id == actors.person["owner"]))
        roles = db.scalar(
            select(func.count()).select_from(UserRoleScope).where(UserRoleScope.user_id == user.id)
        )
    body = departure_body(actors)
    result = departure(actors.admin, actors, body)
    assert result.status_code == 200, result.text
    assert result.json()["employment_status"] == "departed"
    assert departure(actors.admin, actors, body).json() == result.json()
    assert departure(actors.admin, actors, {**body, "evidence_note": "改变依据"}).status_code == 409
    assert actors.owner.get("/api/v1/sessions/current").status_code == 401
    assert actors.other.get("/api/v1/sessions/current").status_code == 200
    with SessionLocal() as db:
        user = db.scalar(select(User).where(User.person_id == actors.person["owner"]))
        assert not user.is_active
        assert not db.scalar(
            select(UserSession.id).where(
                UserSession.user_id == user.id, UserSession.revoked_at.is_(None)
            )
        )
        assert db.get(AccessGrant, grant_id).status == "active"
        assert db.get(Account, UUID(row["id"])).primary_person_id == actors.person["owner"]
        assert (
            db.scalar(
                select(func.count())
                .select_from(UserRoleScope)
                .where(UserRoleScope.user_id == user.id)
            )
            == roles
        )
    records = actors.admin.get("/api/v1/employee-lifecycle").json()["items"]
    target = next(item for item in records if item["id"] == str(actors.person["owner"]))
    assert (
        len(target["handover"]["accounts"]) == 1
        and len(target["handover"]["responsibilities"]) == 1
    )
    assert (
        len(target["handover"]["grants"]) == 1
        and target["last_event"]["evidence_note"] == body["evidence_note"]
    )
    correction = departure_body(
        actors, employment_status="active", departed_on=None, evidence_note="核实并纠正员工状态"
    )
    assert departure(actors.admin, actors, correction).status_code == 200
    with SessionLocal() as db:
        assert not db.scalar(select(User.is_active).where(User.person_id == actors.person["owner"]))


@pytest.mark.parametrize(
    "invalid", ["missing_evidence", "future_date", "stale", "employee", "self", "csrf"]
)
def test_departure_requires_evidence_version_authority_and_csrf(actors, invalid):
    person = "admin" if invalid == "self" else "owner"
    body = departure_body(actors, person)
    if invalid == "missing_evidence":
        body["evidence_note"] = "  "
    if invalid == "future_date":
        body["departed_on"] = (date.today() + timedelta(days=1)).isoformat()
    if invalid == "stale":
        body["expected_updated_at"] = datetime(2020, 1, 1, tzinfo=UTC).isoformat()
    client = actors.other if invalid == "employee" else actors.admin
    if invalid == "csrf":
        browser = TestClient(app)
        try:
            browser.cookies.set(
                get_settings().session_cookie_name, actors.admin.headers["X-PM-Session"]
            )
            assert departure(browser, actors, body).status_code == 403
        finally:
            browser.close()
        return
    response = departure(client, actors, body, person)
    assert response.status_code == (
        422
        if invalid in {"missing_evidence", "future_date"}
        else 403
        if invalid == "employee"
        else 409
    ), response.text


def test_sync_cannot_erase_confirmed_departure_and_old_patch_cannot_bypass(actors):
    assert departure(actors.admin, actors, departure_body(actors)).status_code == 200
    with SessionLocal() as db:
        person = db.get(Person, actors.person["owner"])
        profile = DingTalkPersonProfile(
            person_id=person.id, dingtalk_user_id="DEPART-" + uuid4().hex
        )
        db.add(profile)
        db.commit()
        sync = DingTalkDirectorySync(db, MagicMock())
        sync._upsert_person(
            db.get(LegalEntity, person.legal_entity_id),
            db.get(Department, person.department_id),
            {"userid": profile.dingtalk_user_id, "name": person.display_name, "active": True},
            {},
            SyncResult(),
        )
        db.commit()
        assert person.employment_status == "departed"
    assert (
        actors.admin.patch(
            f"/api/v1/people/{actors.person['owner']}", json={"employment_status": "active"}
        ).status_code
        == 409
    )
    assert actors.other.get("/api/v1/employee-lifecycle").status_code == 403


def test_source_unavailable_or_inactive_does_not_infer_departure(actors):
    with SessionLocal() as db:
        db.get(Person, actors.person["owner"]).employment_status = "inactive"
        db.commit()
    records = actors.admin.get("/api/v1/employee-lifecycle").json()["items"]
    target = next(item for item in records if item["id"] == str(actors.person["owner"]))
    assert target["employment_status"] == "inactive" and target["last_event"] is None


def test_verified_identity_retains_platform_and_custodian_handover(actors, company_account):
    with SessionLocal() as db:
        identity = Asset(
            name="已核验身份",
            asset_code="VER-ID-" + uuid4().hex,
            asset_type_id=db.scalar(
                select(AssetType.id).where(AssetType.code == "registration_identity")
            ),
        )
        db.add(identity)
        db.flush()
        db.add_all(
            [
                RegistrationIdentityProfile(
                    asset_id=identity.id,
                    identity_type="email",
                    identifier_value="verified@example.test",
                    identifier_masked="v***@example.test",
                    identifier_fingerprint=uuid4().hex,
                    source_nature="company_owned",
                    verification_status="verified",
                    custodian_person_id=actors.person["owner"],
                ),
                AssetPlatformLink(asset_id=identity.id, platform_id=company_account["platform"]),
            ]
        )
        db.commit()
        identity_id = str(identity.id)
    result = actors.admin.get("/api/v1/account-hierarchy").json()
    platform = next(
        item for item in result["platforms"] if item["id"] == str(company_account["platform"])
    )
    assert any(
        item["asset_id"] == identity_id and item["name"] == "v***@example.test"
        for item in platform["registration_identities"]
    )
    assert identity_id not in {item["asset_id"] for item in platform["pending_identities"]}
    record = next(
        item
        for item in actors.admin.get("/api/v1/employee-lifecycle").json()["items"]
        if item["id"] == str(actors.person["owner"])
    )
    assert any(item["asset_id"] == identity_id for item in record["handover"]["accounts"])


def test_service_account_without_employee_is_not_an_employee_followup(actors, company_account):
    row, body = child(actors, company_account, account_kind="service")
    with SessionLocal() as db:
        db.add(
            AssetResponsibility(
                asset_id=UUID(company_account["asset"]),
                person_id=actors.person["owner"],
                role_type="responsible",
                is_primary=True,
            )
        )
        tenant = db.scalar(
            select(PlatformTenant).where(PlatformTenant.asset_id == UUID(company_account["asset"]))
        )
        tenant.ownership_nature = "personal_owned"
        db.commit()
    assert company_account["asset"] not in {
        item["id"] for item in actors.admin.get("/api/v1/hudu/expirations").json()["items"]
    }
    hierarchy = actors.admin.get("/api/v1/account-hierarchy").json()
    platform = next(
        item for item in hierarchy["platforms"] if item["id"] == str(company_account["platform"])
    )
    assert any(
        item["asset_id"] == company_account["asset"] for item in platform["personal_accounts"]
    )
    assert not platform["company_accounts"]
    assert binding(actors.admin, company_account, row, binding_body(actors, row)).status_code == 200
    assert (
        actors.admin.post(
            f"/api/v1/workspace/platform-accounts/{company_account['asset']}/child-accounts",
            json=body,
        ).status_code
        == 409
    )

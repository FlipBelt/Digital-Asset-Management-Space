"""Privacy regressions on the explicitly opted-in disposable PostgreSQL database."""

import io
from datetime import UTC, datetime, timedelta
from uuid import UUID, uuid4

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select

from app.api.v1 import usage_monitoring
from app.core.access import (
    AccessContext,
    asset_visibility_clause,
    can_govern_asset,
    can_manage_asset,
    get_access_context,
)
from app.db.session import SessionLocal
from app.models import (
    AccessGrant,
    Account,
    Asset,
    AssetAttachment,
    AssetConfirmation,
    AssetRelation,
    AssetResponsibility,
    AssetType,
    ConnectorDefinition,
    Platform,
    PlatformTenant,
    ProviderConnection,
    ServiceInstance,
    ServiceProduct,
    UsageAlertRule,
    UsageNotificationLog,
    UsageNotificationSchedule,
    User,
)
from app.schemas.usage_alerts import UsageNotificationScheduleWrite
from app.services.asset_confirmation import draft_digest
from app.services.usage_connectors import ensure_usage_catalog
from tests.test_asset_center_integration import confirm, draft, membership, prepare

pytest_plugins = ["tests.test_asset_center_integration"]


def context(db, actors, key, role=None, scopes=()):
    person_id = actors.person[key]
    user = db.scalar(select(User).where(User.person_id == person_id))
    assert user is not None
    if role is None:
        return get_access_context(user, db)
    return AccessContext(
        user, person_id, frozenset({role}), frozenset({"asset.write"}), frozenset(scopes)
    )


@pytest.mark.parametrize(
    ("scope", "funding"),
    [
        ("company", "personal"),
        ("team", "company"),
        (None, "department"),
        ("company", "free"),
        ("company", "trial"),
    ],
)
def test_personal_subscription_ignores_shared_scope_grants_and_department_roles(
    actors, scope, funding
):
    asset, _ = membership(actors.owner, actors, funding)
    asset_id = UUID(asset["id"])
    with SessionLocal() as db:
        item = db.get(Asset, asset_id)
        item.sharing_scope = scope
        item.confidentiality = "internal"
        item.owner_department_id = actors.teams[1]
        ordinary = Asset(
            id=uuid4(),
            name="原有公司资源",
            asset_code="PRIV-COMP-" + uuid4().hex,
            asset_type_id=item.asset_type_id,
            legal_entity_id=UUID(actors.entity),
            created_by_person_id=actors.person["admin"],
            sharing_scope="company",
            source_system=None,
            status="active",
        )
        db.add(ordinary)
        db.flush()
        ordinary_id = str(ordinary.id)
        db.add_all(
            [
                AccessGrant(
                    asset_id=asset_id,
                    person_id=actors.person["other"],
                    grant_type="seat",
                    status="active",
                ),
                AssetResponsibility(
                    asset_id=asset_id, person_id=actors.person["other"], role_type="responsible"
                ),
                AssetResponsibility(
                    asset_id=asset_id, person_id=actors.person["stranger"], role_type="user"
                ),
                AssetResponsibility(
                    asset_id=ordinary.id, person_id=actors.person["other"], role_type="responsible"
                ),
            ]
        )
        db.commit()
        for role in ("employee", "department_manager", "auditor"):
            outsider = context(db, actors, "other", role, [actors.teams[1]])
            assert not db.scalar(
                select(Asset.id).where(Asset.id == asset_id, asset_visibility_clause(outsider))
            )
            assert not can_manage_asset(db, outsider, item)
            assert not can_govern_asset(outsider, item)
        for key, role in (
            ("owner", "employee"),
            ("admin", "system_admin"),
            ("other", "asset_manager"),
        ):
            allowed = context(db, actors, key, role)
            assert db.scalar(
                select(Asset.id).where(Asset.id == asset_id, asset_visibility_clause(allowed))
            )
            assert can_manage_asset(db, allowed, item)
            assert can_govern_asset(allowed, item)
    assert actors.owner.get(f"/api/v1/assets/{asset_id}").status_code == 200
    assert actors.admin.get(f"/api/v1/assets/{asset_id}").status_code == 200
    for suffix in ("", "/assignment", "/relationship-view", "/attachments"):
        assert actors.other.get(f"/api/v1/assets/{asset_id}{suffix}").status_code == 404
    response = actors.other.get("/api/v1/assets", params={"keyword": asset["asset_code"]})
    assert response.status_code == 200 and response.json()["pagination"]["total"] == 0
    for client in (actors.other, actors.stranger):
        mine = client.get("/api/v1/workspace/my-assets")
        assert mine.status_code == 200, mine.text
        assert asset["id"] not in {row["id"] for row in mine.json()}
    assert ordinary_id in {
        row["id"] for row in actors.other.get("/api/v1/workspace/my-assets").json()
    }
    changed = actors.other.put(
        f"/api/v1/assets/{asset_id}/assignment",
        json=dict(
            version=asset["version"],
            responsible_person_id=str(actors.person["other"]),
            owner_department_id=str(actors.teams[1]),
            ownership_scope="department",
        ),
    )
    assert changed.status_code in {403, 404}, changed.text


def test_home_excludes_personal_subscriptions_before_count_and_pagination(actors):
    for client in (actors.owner, actors.admin):
        personal, _ = membership(client, actors)
        normal = draft(client, actors)
        page = client.get("/api/v1/space/assets?scope=mine&category=all&page_size=1").json()
        assert page["pagination"]["total"] == 1
        assert [item["id"] for item in page["data"]] == [normal["id"]]
        page2 = client.get("/api/v1/space/assets?scope=mine&category=all&page=2&page_size=1").json()
        assert page2["pagination"]["total"] == 1 and page2["data"] == []
        subscriptions = client.get("/api/v1/space/assets?category=subscriptions").json()["data"]
        assert personal["id"] in {item["id"] for item in subscriptions}
        assert client.get(f"/api/v1/assets/{personal['id']}").status_code == 200


def test_personal_subscription_hidden_in_hierarchy_export_graph_and_attachment_download(actors):
    asset, _ = membership(actors.owner, actors, "company")
    peer = draft(actors.owner, actors)
    result = confirm(actors.owner, peer, prepare(actors.owner, peer, "company"))
    assert result.status_code == 200
    with SessionLocal() as db:
        db.get(Asset, UUID(asset["id"])).sharing_scope = "company"
        platform = Platform(code="private-" + uuid4().hex, name="隔离隐私平台")
        db.add(platform)
        db.flush()
        db.get(ServiceProduct, UUID(actors.product)).platform_id = platform.id
        attachment = AssetAttachment(
            asset_id=UUID(asset["id"]),
            file_name="private-evidence.txt",
            file_path="/never-served-private-fixture.txt",
            content_type="text/plain",
            size_bytes=1,
        )
        db.add_all(
            [
                attachment,
                AssetRelation(
                    source_asset_id=UUID(peer["id"]),
                    target_asset_id=UUID(asset["id"]),
                    relation_type="uses",
                ),
            ]
        )
        db.commit()
        attachment_id = attachment.id
    for client, visible in ((actors.owner, True), (actors.admin, True), (actors.other, False)):
        hierarchy = client.get("/api/v1/account-hierarchy").json()
        personal_ids = {
            row["asset_id"]
            for platform in hierarchy["platforms"]
            for row in platform["personal_accounts"]
        }
        assert (asset["id"] in personal_ids) == visible
        exported = client.get("/api/v1/exports/assets.xlsx")
        assert exported.status_code == 200
        workbook = load_workbook(io.BytesIO(exported.content), read_only=True)
        codes = {row[0] for row in workbook.active.iter_rows(min_row=2, values_only=True)}
        workbook.close()
        assert (asset["asset_code"] in codes) == visible
        graph = client.get(f"/api/v1/assets/{peer['id']}/relationship-view")
        assert graph.status_code == 200
        assert (asset["id"] in graph.text) == visible
    assert (
        actors.other.get(
            f"/api/v1/assets/{asset['id']}/attachments/{attachment_id}/download"
        ).status_code
        == 404
    )
    assert actors.owner.get(f"/api/v1/assets/{asset['id']}/attachments").status_code == 200


def test_asset_map_hides_private_purchase_account_and_dangling_edges(actors):
    private, _ = membership(actors.owner, actors)
    public, public_instance = membership(actors.admin, actors, "company")
    with SessionLocal() as db:
        ordinary = db.get(Asset, UUID(public["id"]))
        ordinary.source_system = None
        ordinary.sharing_scope = "company"
        ordinary.confidentiality = "internal"
        ordinary.ownership_scope = "company"
        platform = Platform(code="private-map-" + uuid4().hex, name="隔离购买平台")
        db.add(platform)
        db.flush()
        instance = db.get(ServiceInstance, UUID(public_instance["id"]))
        instance.purchase_platform_id = platform.id
        # Historical pointers can predate the current target-type checks.
        instance.purchase_tenant_asset_id = UUID(private["id"])
        db.commit()
    for client, visible in ((actors.owner, True), (actors.admin, True), (actors.other, False)):
        response = client.get("/api/v1/asset-map")
        assert response.status_code == 200, response.text
        graph = response.json()
        node_ids = {row["id"] for row in graph["nodes"]}
        assert f"asset:{public['id']}" in node_ids
        assert (private["id"] in response.text) == visible
        assert all(
            edge["source"] in node_ids and edge["target"] in node_ids for edge in graph["edges"]
        )
        if visible:
            assert any(
                edge["source"] == f"asset:{private['id']}"
                and edge["target"] == f"asset:{public['id']}"
                for edge in graph["edges"]
            )


def test_layer_record_account_counts_match_visible_children(actors):
    private, _ = membership(actors.owner, actors)
    with SessionLocal() as db:
        platform = Platform(code="private-count-" + uuid4().hex, name="隔离计数平台")
        parent = Asset(
            name="可见公司账号",
            asset_code="PRIV-PARENT-" + uuid4().hex,
            asset_type_id=db.scalar(
                select(AssetType.id).where(AssetType.code == "platform_tenant")
            ),
            legal_entity_id=UUID(actors.entity),
            created_by_person_id=actors.person["admin"],
            source_system=None,
            sharing_scope="company",
            ownership_scope="company",
            status="active",
        )
        child = Asset(
            name="可见公司子账号",
            asset_code="PRIV-CHILD-" + uuid4().hex,
            asset_type_id=db.scalar(
                select(AssetType.id).where(AssetType.code == "platform_account")
            ),
            legal_entity_id=UUID(actors.entity),
            created_by_person_id=actors.person["admin"],
            source_system=None,
            sharing_scope="company",
            ownership_scope="company",
            status="active",
        )
        db.add_all([platform, parent, child])
        db.flush()
        tenant = PlatformTenant(
            asset_id=parent.id, platform_id=platform.id, legal_entity_id=UUID(actors.entity)
        )
        db.add(tenant)
        db.flush()
        for asset_id, identifier in (
            (UUID(private["id"]), "private-history@example.test"),
            (child.id, "company-history@example.test"),
        ):
            # Legacy mixed records still must respect the linked asset's privacy.
            db.add(
                Account(
                    asset_id=asset_id,
                    platform_tenant_id=tenant.id,
                    login_identifier=identifier,
                    normalized_login_identifier=identifier,
                    account_type="company_subaccount",
                    registration_identity_type="company",
                )
            )
        db.commit()
        parent_id, child_id = str(parent.id), str(child.id)
    for client, visible in ((actors.owner, True), (actors.admin, True), (actors.other, False)):
        records = client.get("/api/v1/workspace/layer-records?layer=4")
        assert records.status_code == 200, records.text
        record = next(row for row in records.json() if row["asset_id"] == parent_id)
        children = client.get(f"/api/v1/workspace/platform-accounts/{parent_id}/child-accounts")
        assert children.status_code == 200, children.text
        child_ids = {row["asset_id"] for row in children.json()}
        assert child_id in child_ids
        assert (private["id"] in child_ids) == visible
        assert record["account_count"] == len(child_ids) == (2 if visible else 1)


@pytest.mark.parametrize("funding", ["company", "free"])
def test_old_non_private_confirmation_is_rejected_even_for_non_personal_funding(actors, funding):
    asset, _ = membership(actors.owner, actors, funding)
    with SessionLocal() as db:
        item = db.get(Asset, UUID(asset["id"]))
        item.status = "draft"
        item.owner_department_id = actors.teams[0]
        db.flush()
        owner = db.scalar(select(User).where(User.person_id == actors.person["owner"]))
        row = AssetConfirmation(
            asset_id=item.id,
            user_id=owner.id,
            asset_version=item.version,
            outcome_version=1,
            content_digest=draft_digest(db, item, "company"),
            sharing_scope="company",
            expires_at=datetime.now(UTC) + timedelta(minutes=30),
        )
        db.add(row)
        db.commit()
        receipt = dict(
            id=str(row.id), asset_version=row.asset_version, content_digest=row.content_digest
        )
    for scope in ("team", "company"):
        response = actors.owner.post(
            f"/api/v1/assets/{asset['id']}/confirmation",
            json=dict(request_id=str(uuid4()), version=asset["version"], sharing_scope=scope),
        )
        assert response.status_code == 422, response.text
    assert confirm(actors.owner, asset, receipt).status_code == 422
    with SessionLocal() as db:
        assert db.get(AssetConfirmation, UUID(receipt["id"])).consumed_at is None
        assert db.get(Asset, UUID(asset["id"])).sharing_scope == "private"
    private_receipt = prepare(actors.owner, asset)
    assert confirm(actors.owner, asset, private_receipt).status_code == 200


def test_personal_usage_connection_monitoring_and_binding_do_not_bypass_privacy(
    actors, monkeypatch
):
    asset, instance = membership(actors.owner, actors, "free")
    with SessionLocal() as db:
        db.get(Asset, UUID(asset["id"])).sharing_scope = "company"
        ensure_usage_catalog(db)
        definition = db.scalar(
            select(ConnectorDefinition).where(ConnectorDefinition.code == "deepseek")
        )
        private = ProviderConnection(
            connector_definition_id=definition.id,
            legal_entity_id=UUID(actors.entity),
            service_instance_id=UUID(instance["id"]),
            name="私有订阅用量",
            configuration={},
        )
        ordinary = ProviderConnection(
            connector_definition_id=definition.id,
            legal_entity_id=UUID(actors.entity),
            name="原有公司用量连接",
            configuration={},
        )
        db.add_all([private, ordinary])
        db.flush()
        rule = UsageAlertRule(
            provider_connection_id=private.id,
            name="私有额度告警",
            metric_key="balance",
            threshold=10,
            recipient_person_ids=[str(actors.person["owner"])],
        )
        schedule = UsageNotificationSchedule(
            name="含私有连接的日报",
            connection_ids=[str(private.id), str(ordinary.id)],
            recipient_person_ids=[str(actors.person["owner"])],
            time_of_day="09:00",
        )
        db.add_all([rule, schedule])
        db.flush()
        logs = [
            UsageNotificationLog(
                event_type="threshold_breach",
                provider_connection_id=private.id,
                status="sent",
                message_summary="私有直接记录",
            ),
            UsageNotificationLog(
                event_type="threshold_breach",
                usage_alert_rule_id=rule.id,
                status="sent",
                message_summary="私有规则记录",
            ),
            UsageNotificationLog(
                event_type="daily_digest",
                usage_notification_schedule_id=schedule.id,
                status="sent",
                message_summary="私有日报记录",
            ),
        ]
        db.add_all(logs)
        db.commit()
        private_id, ordinary_id, rule_id = private.id, ordinary.id, rule.id
        schedule_id, log_ids = schedule.id, {str(log.id) for log in logs}
        for key, visible in (("owner", True), ("admin", True), ("other", False)):
            schedules = usage_monitoring.list_schedules(db, context(db, actors, key))
            assert (schedule_id in {row["id"] for row in schedules}) == visible
        with pytest.raises(HTTPException) as rejected:
            usage_monitoring.update_schedule(
                schedule_id,
                UsageNotificationScheduleWrite.model_construct(
                    name="试图移除私有连接",
                    connection_ids=[ordinary_id],
                    recipient_person_ids=[actors.person["other"]],
                    time_of_day="09:00",
                ),
                db,
                context(db, actors, "other"),
            )
        assert rejected.value.status_code == 404
    for client, visible in ((actors.owner, True), (actors.admin, True), (actors.other, False)):
        connections = client.get("/api/v1/usage-connections").json()
        ids = {row["id"] for row in connections}
        assert (str(private_id) in ids) == visible and str(ordinary_id) in ids
        overview = client.get("/api/v1/usage-monitoring/overview").json()["connections"]
        assert (str(private_id) in {row["id"] for row in overview}) == visible
        rules = client.get("/api/v1/usage-monitoring/rules").json()
        assert (str(rule_id) in {row["id"] for row in rules}) == visible
        logs = client.get("/api/v1/usage-monitoring/logs").json()
        assert bool(log_ids.intersection(row["id"] for row in logs)) == visible
    for suffix in ("", "/test", "/sync", "/usage-events"):
        response = (
            actors.other.get(f"/api/v1/usage-connections/{private_id}")
            if not suffix
            else actors.other.post(f"/api/v1/usage-connections/{private_id}{suffix}", json={})
        )
        assert response.status_code == 404, response.text
    assert (
        actors.other.patch(
            f"/api/v1/usage-connections/{private_id}", json={"service_instance_id": None}
        ).status_code
        == 404
    )
    assert (
        actors.other.patch(
            f"/api/v1/usage-connections/{ordinary_id}", json={"service_instance_id": instance["id"]}
        ).status_code
        == 404
    )
    assert (
        actors.other.post(
            "/api/v1/usage-connections",
            json=dict(
                platform_code="deepseek",
                legal_entity_id=actors.entity,
                service_instance_id=instance["id"],
                name="非法关联",
                api_key="test-key",
            ),
        ).status_code
        == 404
    )
    rule_payload = dict(
        provider_connection_id=str(ordinary_id),
        name="试图移除私有关联",
        threshold=10,
        recipient_person_ids=[str(actors.person["other"])],
    )
    assert (
        actors.other.put(f"/api/v1/usage-monitoring/rules/{rule_id}", json=rule_payload).status_code
        == 404
    )
    monkeypatch.setattr(
        usage_monitoring, "run_monitoring_cycle", lambda db: pytest.fail("must not run")
    )
    assert actors.other.post("/api/v1/usage-monitoring/run").status_code == 403

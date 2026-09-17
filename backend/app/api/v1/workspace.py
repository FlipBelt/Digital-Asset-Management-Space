import csv
import hashlib
import io
import json
import re
from datetime import UTC, datetime
from decimal import Decimal
from html.parser import HTMLParser
from typing import Annotated
from uuid import UUID

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.access import (
    AccessContext,
    asset_visibility_clause,
    get_access_context,
    require_asset_visible,
    require_asset_write,
)
from app.core.config import get_settings
from app.db.session import get_db
from app.models import (
    AccessGrant,
    Account,
    Asset,
    AssetIdentifier,
    AssetPlatformLink,
    AssetRelation,
    AssetResponsibility,
    AssetType,
    AuditLog,
    Department,
    ImportAnalysis,
    ImportBatch,
    ImportProposedObject,
    ImportProposedRelation,
    InternalSystemProfile,
    LegalEntity,
    Person,
    Platform,
    PlatformAccountRegistrationIdentity,
    PlatformTenant,
    Provider,
    RegistrationIdentityProfile,
    ResourceProfile,
    ServiceInstance,
    ServiceProduct,
    SourceImportRecord,
)
from app.schemas.assets import AssetRead
from app.schemas.inventory import AccountRead
from app.schemas.workspace import (
    AccessGrantCreate,
    AssetMapRead,
    BossPilotImportRequest,
    BossPilotImportResult,
    BossPilotRollbackResult,
    CompanyPlatformAccountCreate,
    FlexibleImportPreview,
    FlexibleImportStage,
    FlexibleImportStageResult,
    ImportAccountCandidateRead,
    ImportBatchRead,
    ImportCandidate,
    ImportCommitRequest,
    ImportCommitResult,
    ImportDryRunResult,
    ImportPlanRead,
    ImportProposalObjectPatch,
    ImportProposalObjectRead,
    ImportProposalRelationPatch,
    ImportProposalRelationRead,
    IntakeLinkRead,
    IntakeNextActionRead,
    IntakeResultRead,
    LayerRecordRead,
    LegacyHierarchyNormalizeResult,
    MapEdge,
    MapNode,
    PlatformAccountChildCreate,
    PlatformAccountContextRead,
    RegistrationIdentityCreate,
    RegistrationIdentityRead,
    ResourceCreate,
)
from app.services.aliyun_colleague_import import (
    commit_aliyun_colleague_import,
    reclassify_aliyun_colleague_as_l4_accounts,
    rollback_aliyun_colleague_import,
)
from app.services.assets import allocate_asset_code, ensure_internal_identifier
from app.services.import_commit import commit_reviewed_import_batch
from app.services.import_planning import (
    AI_ANALYZER_VERSION,
    AIImportError,
    dry_run_import_plan,
    enhance_import_records_with_ai,
    rebuild_import_plan,
)

router = APIRouter(tags=["workspace"])


def require_asset_type(db: Session, code: str) -> AssetType:
    item = db.scalar(select(AssetType).where(AssetType.code == code))
    if item is None:
        raise HTTPException(status_code=422, detail=f"缺少资产类型：{code}")
    return item


def require_legal_entity(db: Session, legal_entity_id: UUID | None) -> LegalEntity | None:
    if legal_entity_id is None:
        return None
    item = db.get(LegalEntity, legal_entity_id)
    if item is None or item.archived_at is not None:
        raise HTTPException(status_code=422, detail="所选公司主体不存在")
    return item


def create_asset(
    db: Session,
    *,
    name: str,
    asset_type_id: UUID,
    legal_entity_id: UUID | None,
    status_value: str,
    description: str | None,
    owner_department_id: UUID | None = None,
    criticality: str = "normal",
    ownership_scope: str = "pending",
    source_type: str = "manual",
    created_by_person_id: UUID | None = None,
    review_status: str = "pending_review",
) -> Asset:
    item = Asset(
        asset_code=allocate_asset_code(db, legal_entity_id, asset_type_id),
        name=name,
        asset_type_id=asset_type_id,
        legal_entity_id=legal_entity_id,
        owner_department_id=owner_department_id,
        ownership_scope=ownership_scope,
        status=status_value,
        criticality=criticality,
        confidentiality="internal",
        source_type=source_type,
        description=description,
        created_by_person_id=created_by_person_id,
        review_status=review_status,
    )
    db.add(item)
    db.flush()
    ensure_internal_identifier(db, item)
    return item


def add_discovery_audit(
    db: Session, access: AccessContext, asset: Asset, discovery_kind: str
) -> None:
    db.add(
        AuditLog(
            actor_user_id=access.user.id,
            action="asset.discovery.create",
            object_type="asset",
            object_id=asset.id,
            after_data={
                "name": asset.name,
                "status": asset.status,
                "review_status": asset.review_status,
                "discovery_kind": discovery_kind,
                "owner_department_id": (
                    str(asset.owner_department_id) if asset.owner_department_id else None
                ),
            },
            request_id="workspace-intake",
        )
    )


def infer_identity_type(value: str) -> str:
    normalized = value.strip()
    if re.fullmatch(r"1\d{10}", normalized):
        return "phone"
    if re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", normalized):
        return "email"
    if normalized.lower().startswith(("wx:", "wechat:")):
        return "wechat"
    return "other"


@router.get("/workspace/registration-identities", response_model=list[RegistrationIdentityRead])
def list_registration_identities(
    db: Session = Depends(get_db),
    access: AccessContext = Depends(get_access_context),
) -> list[RegistrationIdentityRead]:
    rows = db.execute(
        select(RegistrationIdentityProfile, Asset)
        .join(Asset, Asset.id == RegistrationIdentityProfile.asset_id)
        .where(
            RegistrationIdentityProfile.archived_at.is_(None),
            Asset.archived_at.is_(None),
            asset_visibility_clause(access),
        )
        .order_by(Asset.name)
    ).all()
    return [
        RegistrationIdentityRead(
            id=profile.id,
            asset_id=asset.id,
            asset_code=asset.asset_code,
            legal_entity_id=asset.legal_entity_id,
            name=asset.name,
            identity_type=profile.identity_type,
            identifier=profile.identifier_value,
            source_nature=profile.source_nature,
            custodian_person_id=profile.custodian_person_id,
            verification_status=profile.verification_status,
        )
        for profile, asset in rows
    ]


@router.get("/workspace/my-assets", response_model=list[AssetRead])
def list_my_assets(
    db: Session = Depends(get_db),
    access: AccessContext = Depends(get_access_context),
) -> list[AssetRead]:
    """Return the current person's assets in one filtered database query."""
    if access.person_id is None:
        return []
    rows = db.scalars(
        select(Asset)
        .join(AssetResponsibility, AssetResponsibility.asset_id == Asset.id)
        .where(
            Asset.archived_at.is_(None),
            AssetResponsibility.person_id == access.person_id,
            AssetResponsibility.role_type.in_(["responsible", "user"]),
            AssetResponsibility.archived_at.is_(None),
        )
        .distinct()
        .order_by(Asset.updated_at.desc(), Asset.id)
    )
    return [AssetRead.model_validate(item) for item in rows]


@router.get("/workspace/layer-records", response_model=list[LayerRecordRead])
def list_layer_records(
    layer: int = 6,
    db: Session = Depends(get_db),
    access: AccessContext = Depends(get_access_context),
) -> list[LayerRecordRead]:
    """A deliberate six-layer browser; each layer returns only real records."""
    if layer not in {1, 2, 3, 4, 5, 6}:
        raise HTTPException(status_code=422, detail="层级必须为 1 至 6")

    visible_assets = list(
        db.scalars(
            select(Asset).where(Asset.archived_at.is_(None), asset_visibility_clause(access))
        )
    )
    visible_asset_ids = {item.id for item in visible_assets}
    entities = {
        item.id: item
        for item in db.scalars(select(LegalEntity).where(LegalEntity.archived_at.is_(None)))
    }
    asset_types = {
        item.id: item
        for item in db.scalars(select(AssetType).where(AssetType.archived_at.is_(None)))
    }
    tenants = list(
        db.scalars(
            select(PlatformTenant).where(
                PlatformTenant.archived_at.is_(None),
                PlatformTenant.asset_id.in_(visible_asset_ids),
            )
        )
    ) if visible_asset_ids else []
    account_counts = {
        tenant_id: count
        for tenant_id, count in db.execute(
            select(Account.platform_tenant_id, func.count(Account.id))
            .where(
                Account.platform_tenant_id.in_([item.id for item in tenants]),
                Account.archived_at.is_(None),
            )
            .group_by(Account.platform_tenant_id)
        )
    } if tenants else {}
    service_instances = list(
        db.scalars(
            select(ServiceInstance).where(
                ServiceInstance.archived_at.is_(None),
                ServiceInstance.asset_id.in_(visible_asset_ids),
            )
        )
    ) if visible_asset_ids else []
    # L1 and L3 are independent directory objects.  They must remain visible
    # even when the evidence does not justify an L4/L6 relation yet; otherwise
    # a correctly imported, intentionally unlinked entity looks as if it was
    # never stored.
    platforms = {
        item.id: item
        for item in db.scalars(select(Platform).where(Platform.archived_at.is_(None)))
    }
    direct_platform_links = (
        list(
            db.scalars(
                select(AssetPlatformLink).where(
                    AssetPlatformLink.asset_id.in_(visible_asset_ids),
                    AssetPlatformLink.archived_at.is_(None),
                )
            )
        )
        if visible_asset_ids
        else []
    )
    platform_names_by_asset_id: dict[UUID, list[str]] = {}
    for link in direct_platform_links:
        platform_name = platforms.get(link.platform_id).name if link.platform_id in platforms else ""
        if platform_name and platform_name not in platform_names_by_asset_id.setdefault(link.asset_id, []):
            platform_names_by_asset_id[link.asset_id].append(platform_name)

    def platform_names(asset_id: UUID) -> str | None:
        names = platform_names_by_asset_id.get(asset_id, [])
        return "、".join(names) or None

    if layer == 1:
        return [
            LayerRecordRead(
                id=f"entity:{item.id}",
                layer=1,
                name=item.name,
                object_type="公司主体",
                status=item.status,
                updated_at=item.updated_at,
            )
            for item in entities.values()
        ]
    if layer == 2:
        profiles = list(
            db.scalars(
                select(RegistrationIdentityProfile).where(
                    RegistrationIdentityProfile.archived_at.is_(None),
                    RegistrationIdentityProfile.asset_id.in_(visible_asset_ids),
                )
            )
        ) if visible_asset_ids else []
        assets_by_id = {asset.id: asset for asset in visible_assets}
        return [
            LayerRecordRead(
                id=f"identity:{profile.id}",
                layer=2,
                name=profile.identifier_value,
                object_type=identity_type_label(profile.identity_type),
                legal_entity_name=(
                    entities[assets_by_id[profile.asset_id].legal_entity_id].name
                    if assets_by_id[profile.asset_id].legal_entity_id in entities
                    else None
                ),
                platform_name=platform_names(profile.asset_id),
                ownership_nature=profile.source_nature,
                status=profile.verification_status,
                asset_id=profile.asset_id,
                updated_at=profile.updated_at,
            )
            for profile in profiles
            if profile.asset_id in assets_by_id
        ]
    if layer == 3:
        return [
            LayerRecordRead(
                id=f"platform:{item.id}",
                layer=3,
                name=item.name,
                object_type="平台",
                category=item.category,
                status=item.review_status,
                updated_at=item.updated_at,
            )
            for item in platforms.values()
        ]
    if layer == 4:
        assets_by_id = {asset.id: asset for asset in visible_assets}
        return [
            LayerRecordRead(
                id=f"platform-account:{tenant.id}",
                layer=4,
                name=assets_by_id[tenant.asset_id].name,
                object_type="公司平台账号",
                legal_entity_name=(
                    entities[tenant.legal_entity_id].name
                    if tenant.legal_entity_id in entities
                    else None
                ),
                platform_name=(
                    platforms[tenant.platform_id].name
                    if tenant.platform_id in platforms
                    else None
                ),
                ownership_nature=tenant.ownership_nature,
                account_count=account_counts.get(tenant.id, 0),
                status=tenant.verification_status,
                asset_id=tenant.asset_id,
                updated_at=tenant.updated_at,
            )
            for tenant in tenants
            if tenant.asset_id in assets_by_id
        ]
    if layer == 5:
        people = {
            item.id: item
            for item in db.scalars(select(Person).where(Person.archived_at.is_(None)))
        }
        assets_by_id = {asset.id: asset for asset in visible_assets}
        grants = list(
            db.scalars(
                select(AccessGrant).where(
                    AccessGrant.archived_at.is_(None),
                    AccessGrant.asset_id.in_(visible_asset_ids),
                )
            )
        ) if visible_asset_ids else []

        def grant_name(grant: AccessGrant) -> str:
            person_name = (
                people[grant.person_id].display_name
                if grant.person_id in people
                else "待确认人员"
            )
            target_name = assets_by_id[grant.asset_id].name if grant.asset_id in assets_by_id else "待确认资产"
            return f"{person_name} → {target_name}"

        return [
            LayerRecordRead(
                id=f"grant:{grant.id}",
                layer=5,
                name=grant_name(grant),
                object_type="人员授权",
                legal_entity_name=(
                    entities[assets_by_id[grant.asset_id].legal_entity_id].name
                    if assets_by_id[grant.asset_id].legal_entity_id in entities
                    else None
                ),
                status=grant.status,
                asset_id=grant.asset_id,
                updated_at=grant.updated_at,
            )
            for grant in grants
        ]

    layer_type_codes = {"registration_identity", "platform_tenant", "platform_account"}
    service_by_asset_id = {item.asset_id: item for item in service_instances}
    return [
        LayerRecordRead(
            id=f"asset:{asset.id}",
            layer=6,
            name=asset.name,
            object_type=(
                asset_types[asset.asset_type_id].name
                if asset.asset_type_id in asset_types
                else "服务/资源"
            ),
            legal_entity_name=(
                entities[asset.legal_entity_id].name
                if asset.legal_entity_id in entities
                else None
            ),
            platform_name=(
                platforms.get(service_by_asset_id[asset.id].purchase_platform_id).name
                if asset.id in service_by_asset_id
                and service_by_asset_id[asset.id].purchase_platform_id in platforms
                else None
            ) or platform_names(asset.id),
            status=asset.status,
            asset_id=asset.id,
            updated_at=asset.updated_at,
        )
        for asset in visible_assets
        if asset_types.get(asset.asset_type_id) is None
        or asset_types[asset.asset_type_id].code not in layer_type_codes
    ]


@router.post(
    "/workspace/registration-identities",
    response_model=RegistrationIdentityRead,
    status_code=status.HTTP_201_CREATED,
)
def create_registration_identity(
    payload: RegistrationIdentityCreate,
    db: Session = Depends(get_db),
    access: AccessContext = Depends(get_access_context),
) -> RegistrationIdentityRead:
    require_legal_entity(db, payload.legal_entity_id)
    if payload.custodian_person_id:
        custodian = db.get(Person, payload.custodian_person_id)
        if (
            custodian is None
            or custodian.archived_at is not None
            or (
                payload.legal_entity_id is not None
                and custodian.legal_entity_id != payload.legal_entity_id
            )
        ):
            raise HTTPException(status_code=422, detail="当前保管人不属于所选公司")
    value = payload.identifier.strip()
    fingerprint = hashlib.sha256(value.lower().encode("utf-8")).hexdigest()
    existing = db.scalar(
        select(RegistrationIdentityProfile).where(
            RegistrationIdentityProfile.identifier_fingerprint == fingerprint,
            RegistrationIdentityProfile.archived_at.is_(None),
        )
    )
    if existing is not None:
        raise HTTPException(status_code=409, detail="该注册身份已经登记")
    identity_type = payload.identity_type or infer_identity_type(value)
    asset_type = require_asset_type(db, "registration_identity")
    asset = create_asset(
        db,
        name=f"{identity_type_label(identity_type)} · {value}",
        asset_type_id=asset_type.id,
        legal_entity_id=payload.legal_entity_id,
        status_value="draft",
        description=payload.note,
        created_by_person_id=access.person_id,
    )
    profile = RegistrationIdentityProfile(
        asset_id=asset.id,
        identity_type=identity_type,
        identifier_value=value,
        identifier_masked=value,
        identifier_fingerprint=fingerprint,
        source_nature=payload.source_nature,
        custodian_person_id=payload.custodian_person_id,
        verification_status=payload.verification_status,
    )
    db.add(profile)
    if payload.platform_id:
        platform = db.get(Platform, payload.platform_id)
        if platform is None or platform.archived_at is not None:
            raise HTTPException(status_code=422, detail="所选平台不存在")
        db.add(
            AssetPlatformLink(
                asset_id=asset.id,
                platform_id=platform.id,
                relation_type="registered_on",
                source_type="workspace_form",
                review_status="pending_review",
                note="注册身份与平台的显式关联；不自动创建L4。",
            )
        )
    add_discovery_audit(db, access, asset, "registration_identity")
    db.commit()
    db.refresh(profile)
    return RegistrationIdentityRead(
        id=profile.id,
        asset_id=asset.id,
        asset_code=asset.asset_code,
        legal_entity_id=asset.legal_entity_id,
        name=asset.name,
        identity_type=profile.identity_type,
        identifier=profile.identifier_value,
        source_nature=profile.source_nature,
        custodian_person_id=profile.custodian_person_id,
        verification_status=profile.verification_status,
    )


def identity_type_label(identity_type: str) -> str:
    return {"phone": "手机号", "email": "邮箱", "wechat": "微信身份"}.get(identity_type, "注册身份")


@router.get(
    "/workspace/platform-accounts/by-asset/{asset_id}",
    response_model=PlatformAccountContextRead,
)
def get_platform_account_context(
    asset_id: UUID,
    db: Session = Depends(get_db),
    access: AccessContext = Depends(get_access_context),
) -> PlatformAccountContextRead:
    asset = db.get(Asset, asset_id)
    if asset is None or asset.archived_at is not None:
        raise HTTPException(status_code=404, detail="平台账号不存在")
    require_asset_visible(db, access, asset)
    tenant = db.scalar(
        select(PlatformTenant).where(
            PlatformTenant.asset_id == asset_id, PlatformTenant.archived_at.is_(None)
        )
    )
    if tenant is None:
        raise HTTPException(status_code=404, detail="该资产不是平台账号")
    platform = db.get(Platform, tenant.platform_id)
    if platform is None:
        raise HTTPException(status_code=422, detail="平台账号缺少所属平台")
    identity_rows = db.execute(
        select(PlatformAccountRegistrationIdentity, RegistrationIdentityProfile, Asset)
        .join(
            RegistrationIdentityProfile,
            RegistrationIdentityProfile.asset_id
            == PlatformAccountRegistrationIdentity.registration_identity_asset_id,
        )
        .join(Asset, Asset.id == RegistrationIdentityProfile.asset_id)
        .where(
            PlatformAccountRegistrationIdentity.platform_tenant_id == tenant.id,
            PlatformAccountRegistrationIdentity.status == "active",
            RegistrationIdentityProfile.archived_at.is_(None),
            Asset.archived_at.is_(None),
        )
        .order_by(PlatformAccountRegistrationIdentity.created_at)
    ).all()
    return PlatformAccountContextRead(
        asset_id=asset.id,
        platform_id=platform.id,
        platform_name=platform.name,
        platform_category="平台",
        platform_website=platform.website,
        tenant_identifier=tenant.tenant_identifier,
        external_identifier_type=tenant.external_identifier_type,
        ownership_nature=tenant.ownership_nature,
        account_scope=tenant.account_scope,
        verification_status=tenant.verification_status,
        evidence_note=tenant.evidence_note,
        registration_identities=[
            RegistrationIdentityRead(
                id=profile.id,
                asset_id=identity_asset.id,
                asset_code=identity_asset.asset_code,
                legal_entity_id=identity_asset.legal_entity_id,
                name=identity_asset.name,
                identity_type=profile.identity_type,
                identifier=profile.identifier_value,
                source_nature=profile.source_nature,
                custodian_person_id=profile.custodian_person_id,
                verification_status=profile.verification_status,
            )
            for _, profile, identity_asset in identity_rows
        ],
    )


def require_platform_tenant_asset(
    db: Session, access: AccessContext, asset_id: UUID
) -> tuple[Asset, PlatformTenant]:
    asset = db.get(Asset, asset_id)
    if asset is None or asset.archived_at is not None:
        raise HTTPException(status_code=404, detail="公司平台账号不存在")
    require_asset_visible(db, access, asset)
    tenant = db.scalar(
        select(PlatformTenant).where(
            PlatformTenant.asset_id == asset.id,
            PlatformTenant.archived_at.is_(None),
        )
    )
    if tenant is None:
        raise HTTPException(status_code=422, detail="当前记录不是L4公司平台账号")
    return asset, tenant


@router.get(
    "/workspace/platform-accounts/{asset_id}/child-accounts",
    response_model=list[AccountRead],
)
def list_platform_account_children(
    asset_id: UUID,
    db: Session = Depends(get_db),
    access: AccessContext = Depends(get_access_context),
) -> list[Account]:
    """Account and seat details live under L4 and are never a seventh layer."""
    _, tenant = require_platform_tenant_asset(db, access, asset_id)
    return list(
        db.scalars(
            select(Account)
            .where(
                Account.platform_tenant_id == tenant.id,
                Account.archived_at.is_(None),
            )
            .order_by(Account.created_at)
        )
    )


@router.post(
    "/workspace/platform-accounts/{asset_id}/child-accounts",
    response_model=AccountRead,
    status_code=status.HTTP_201_CREATED,
)
def create_platform_account_child(
    asset_id: UUID,
    payload: PlatformAccountChildCreate,
    db: Session = Depends(get_db),
    access: AccessContext = Depends(require_asset_write),
) -> Account:
    parent_asset, tenant = require_platform_tenant_asset(db, access, asset_id)
    normalized = payload.login_identifier.strip().lower()
    duplicate = db.scalar(
        select(Account).where(
            Account.platform_tenant_id == tenant.id,
            Account.normalized_login_identifier == normalized,
            Account.archived_at.is_(None),
        )
    )
    if duplicate is not None:
        raise HTTPException(status_code=409, detail="该公司平台账号下已存在相同登录标识")
    if payload.parent_account_id:
        parent_account = db.get(Account, payload.parent_account_id)
        if (
            parent_account is None
            or parent_account.archived_at is not None
            or parent_account.platform_tenant_id != tenant.id
        ):
            raise HTTPException(status_code=422, detail="上级账号不属于当前公司平台账号")

    account_type = require_asset_type(db, "platform_account")
    account_asset = create_asset(
        db,
        name=f"{parent_asset.name} · {payload.display_name}",
        asset_type_id=account_type.id,
        legal_entity_id=tenant.legal_entity_id,
        status_value="draft",
        description=(
            "L4公司平台账号下的登录账号/席位明细；不属于六层独立对象，不保存密码。"
            + (f"\n来源说明：{payload.note}" if payload.note else "")
        ),
        criticality="important" if payload.privilege_level in {"admin", "root"} else "normal",
        source_type="manual",
        created_by_person_id=access.person_id,
        review_status="pending_review",
    )
    item = Account(
        asset_id=account_asset.id,
        platform_tenant_id=tenant.id,
        login_identifier=payload.login_identifier.strip(),
        normalized_login_identifier=normalized,
        account_type=payload.account_type,
        registration_identity_type="account_detail",
        mfa_status=payload.mfa_status,
        privilege_level=payload.privilege_level,
        parent_account_id=payload.parent_account_id,
        account_kind=payload.account_kind,
        login_method=payload.login_method,
        account_role=payload.account_role,
    )
    db.add(item)
    db.add(
        AuditLog(
            actor_user_id=access.user.id,
            action="platform_account.child_account.create",
            object_type="account",
            object_id=item.id,
            after_data={
                "platform_tenant_id": str(tenant.id),
                "login_identifier": item.login_identifier,
                "account_kind": item.account_kind,
                "account_role": item.account_role,
            },
            request_id="workspace-platform-account-child",
        )
    )
    db.commit()
    db.refresh(item)
    return item


@router.post(
    "/workspace/platform-accounts",
    response_model=IntakeResultRead,
    status_code=status.HTTP_201_CREATED,
)
def create_company_platform_account(
    payload: CompanyPlatformAccountCreate,
    db: Session = Depends(get_db),
    access: AccessContext = Depends(get_access_context),
) -> IntakeResultRead:
    require_legal_entity(db, payload.legal_entity_id)
    platform = db.get(Platform, payload.platform_id)
    if platform is None:
        raise HTTPException(status_code=422, detail="所选平台不存在")
    if payload.ownership_scope not in {"company", "department", "pending"}:
        raise HTTPException(status_code=422, detail="归属范围无效")
    if payload.ownership_scope == "department" and not payload.owner_department_id:
        raise HTTPException(status_code=422, detail="部门级账号必须选择归属部门")
    if payload.ownership_scope != "department" and payload.owner_department_id:
        raise HTTPException(status_code=422, detail="只有部门级账号可以设置归属部门")
    if payload.owner_department_id:
        department = db.get(Department, payload.owner_department_id)
        if (
            department is None
            or department.archived_at is not None
            or department.legal_entity_id != payload.legal_entity_id
        ):
            raise HTTPException(status_code=422, detail="归属部门不属于所选公司")
    asset_type = require_asset_type(db, "platform_tenant")
    asset = create_asset(
        db,
        name=payload.internal_name,
        asset_type_id=asset_type.id,
        legal_entity_id=payload.legal_entity_id,
        owner_department_id=payload.owner_department_id,
        status_value="draft",
        description=payload.description,
        created_by_person_id=access.person_id,
    )
    tenant = PlatformTenant(
        asset_id=asset.id,
        platform_id=payload.platform_id,
        legal_entity_id=payload.legal_entity_id,
        tenant_identifier=payload.external_identifier_value,
        external_identifier_type=payload.external_identifier_type,
        ownership_nature=payload.ownership_nature,
        account_scope=payload.account_scope,
        verification_status="pending",
        evidence_note=payload.evidence_note,
    )
    db.add(tenant)
    db.flush()
    identity_asset: Asset | None = None
    identity_profile: RegistrationIdentityProfile | None = None
    if payload.registration_identity_asset_id:
        identity_asset = db.get(Asset, payload.registration_identity_asset_id)
        if identity_asset is None:
            raise HTTPException(status_code=422, detail="所选注册身份不存在")
        require_asset_visible(db, access, identity_asset)
        if identity_asset.legal_entity_id != payload.legal_entity_id:
            raise HTTPException(status_code=422, detail="注册身份与平台账号必须属于同一公司")
        identity_profile = db.scalar(
            select(RegistrationIdentityProfile).where(
                RegistrationIdentityProfile.asset_id == identity_asset.id,
                RegistrationIdentityProfile.archived_at.is_(None),
            )
        )
        db.add(
            PlatformAccountRegistrationIdentity(
                platform_tenant_id=tenant.id,
                registration_identity_asset_id=payload.registration_identity_asset_id,
                role="primary",
                status="active",
            )
        )
        db.add(
            AssetRelation(
                source_asset_id=asset.id,
                target_asset_id=identity_asset.id,
                relation_type="registered_by",
                source_type="workspace_form",
            )
        )

    responsible_person_id = payload.responsible_person_id
    if responsible_person_id is None and identity_profile is not None:
        responsible_person_id = identity_profile.custodian_person_id
    if responsible_person_id is None:
        responsible_person_id = access.person_id
    proposed_person_ids = {
        person_id
        for person_id in [responsible_person_id, *payload.user_person_ids]
        if person_id is not None
    }
    people = list(
        db.scalars(
            select(Person).where(Person.id.in_(proposed_person_ids), Person.archived_at.is_(None))
        )
    )
    if len({person.id for person in people}) != len(proposed_person_ids) or any(
        person.legal_entity_id != payload.legal_entity_id for person in people
    ):
        raise HTTPException(status_code=422, detail="负责人和使用人员必须属于所选公司")
    if responsible_person_id:
        db.add(
            AssetResponsibility(
                asset_id=asset.id,
                person_id=responsible_person_id,
                role_type="proposed_responsible",
                is_primary=True,
            )
        )
    for person_id in dict.fromkeys(payload.user_person_ids):
        if person_id != responsible_person_id:
            db.add(
                AssetResponsibility(
                    asset_id=asset.id,
                    person_id=person_id,
                    role_type="proposed_user",
                    is_primary=False,
                )
            )
    asset.ownership_scope = payload.ownership_scope
    add_discovery_audit(db, access, asset, "platform_account")
    db.commit()
    links = [
        IntakeLinkRead(kind="platform", label=platform.name, relation="所属平台"),
    ]
    if identity_asset is not None:
        links.append(
            IntakeLinkRead(
                kind="registration_identity",
                label=identity_asset.name,
                relation="主要注册身份",
                asset_id=identity_asset.id,
            )
        )
    if responsible_person_id:
        responsible = next(
            (person for person in people if person.id == responsible_person_id), None
        )
        if responsible:
            links.append(
                IntakeLinkRead(kind="person", label=responsible.display_name, relation="建议负责人")
            )
    completion = 55 + 15 + (15 if identity_asset else 0) + (10 if responsible_person_id else 0)
    if tenant.tenant_identifier:
        completion += 5
    return IntakeResultRead(
        id=tenant.id,
        asset_id=asset.id,
        asset_code=asset.asset_code,
        name=asset.name,
        object_type="platform_account",
        completion_percent=min(100, completion),
        links=links,
        next_actions=[
            IntakeNextActionRead(
                key="confirm_assignment",
                label="确认归属与责任",
                description="系统已带入建议负责人和归属范围，确认后转为正式资产。",
                target=f"/assets/{asset.id}?tab=responsibility",
                required=True,
            ),
            IntakeNextActionRead(
                key="complete_profile",
                label="补充账号资料",
                description="可继续填写 MFA、管理入口等平台账号专属信息。",
                target=f"/assets/{asset.id}?tab=profile",
            ),
            IntakeNextActionRead(
                key="register_resource",
                label="登记这个账号管理的服务",
                description="例如订阅、API、服务器或其他资源。",
                target=f"/intake?mode=resource&account={tenant.id}",
            ),
        ],
    )


@router.post(
    "/workspace/resources",
    response_model=IntakeResultRead,
    status_code=status.HTTP_201_CREATED,
)
def create_resource(
    payload: ResourceCreate,
    db: Session = Depends(get_db),
    access: AccessContext = Depends(get_access_context),
) -> IntakeResultRead:
    selected_tenant: PlatformTenant | None = None
    managed_asset: Asset | None = None
    if payload.managed_under_account_id:
        selected_tenant = db.get(PlatformTenant, payload.managed_under_account_id)
        if selected_tenant is None or selected_tenant.archived_at is not None:
            raise HTTPException(status_code=422, detail="所选公司平台账号不存在")
        managed_asset = db.get(Asset, selected_tenant.asset_id)
        if managed_asset is None or managed_asset.archived_at is not None:
            raise HTTPException(status_code=422, detail="所选公司平台账号的资产记录不存在")
        require_asset_visible(db, access, managed_asset)
        if payload.legal_entity_id is not None and selected_tenant.legal_entity_id != payload.legal_entity_id:
            raise HTTPException(status_code=422, detail="管理账号与资源必须属于同一公司")
        if payload.legal_entity_id is None:
            payload.legal_entity_id = selected_tenant.legal_entity_id
    require_legal_entity(db, payload.legal_entity_id)
    if db.get(AssetType, payload.asset_type_id) is None:
        raise HTTPException(status_code=422, detail="所选资源类型不存在")
    asset = create_asset(
        db,
        name=payload.name,
        asset_type_id=payload.asset_type_id,
        legal_entity_id=payload.legal_entity_id,
        owner_department_id=payload.owner_department_id,
        status_value="draft",
        criticality=payload.criticality,
        description=payload.business_purpose,
        created_by_person_id=access.person_id,
    )
    profile = ResourceProfile(
        asset_id=asset.id,
        resource_family=payload.resource_family,
        managed_under_account_id=payload.managed_under_account_id,
        parent_resource_asset_id=payload.parent_resource_asset_id,
        external_identifier_type=payload.external_identifier_type,
        external_identifier_value=payload.external_identifier_value,
        management_url=payload.management_url,
        verification_status="pending",
    )
    db.add(profile)
    if selected_tenant and managed_asset:
        db.add(
            AssetRelation(
                source_asset_id=managed_asset.id,
                target_asset_id=asset.id,
                relation_type="manages",
                source_type="workspace_form",
            )
        )
    if payload.platform_id:
        platform = db.get(Platform, payload.platform_id)
        if platform is None or platform.archived_at is not None:
            raise HTTPException(status_code=422, detail="所选平台不存在")
        db.add(
            AssetPlatformLink(
                asset_id=asset.id,
                platform_id=platform.id,
                relation_type=payload.platform_relation_type,
                source_type="workspace_form",
                review_status="pending_review",
                note="资源与平台的显式关联；不自动创建L4。",
            )
        )
    if payload.parent_resource_asset_id:
        parent_asset = db.get(Asset, payload.parent_resource_asset_id)
        if parent_asset is None:
            raise HTTPException(status_code=422, detail="所选上级资源不存在")
        require_asset_visible(db, access, parent_asset)
        db.add(
            AssetRelation(
                source_asset_id=payload.parent_resource_asset_id,
                target_asset_id=asset.id,
                relation_type="contains",
                source_type="workspace_form",
            )
        )
    add_proposed_assignment(db, asset, payload)
    add_discovery_audit(db, access, asset, "resource")
    db.commit()
    links: list[IntakeLinkRead] = []
    if payload.managed_under_account_id:
        tenant = db.get(PlatformTenant, payload.managed_under_account_id)
        managed_asset = db.get(Asset, tenant.asset_id) if tenant else None
        if managed_asset:
            links.append(
                IntakeLinkRead(
                    kind="platform_account",
                    label=managed_asset.name,
                    relation="管理来源",
                    asset_id=managed_asset.id,
                )
            )
    if payload.responsible_person_id:
        responsible = db.get(Person, payload.responsible_person_id)
        if responsible:
            links.append(
                IntakeLinkRead(kind="person", label=responsible.display_name, relation="建议负责人")
            )
    completion = 55 + (15 if payload.managed_under_account_id else 0)
    completion += 15 if payload.responsible_person_id else 0
    completion += 10 if payload.external_identifier_value else 0
    return IntakeResultRead(
        id=profile.id,
        asset_id=asset.id,
        asset_code=asset.asset_code,
        name=asset.name,
        object_type="resource",
        completion_percent=min(100, completion),
        links=links,
        next_actions=[
            IntakeNextActionRead(
                key="confirm_assignment",
                label="确认归属与责任",
                description="核对建议归属、负责人和实际使用人员。",
                target=f"/assets/{asset.id}?tab=responsibility",
                required=True,
            ),
            IntakeNextActionRead(
                key="complete_profile",
                label="补充专属资料",
                description="按照资源类型补充实例编号、区域或管理地址。",
                target=f"/assets/{asset.id}?tab=profile",
            ),
            IntakeNextActionRead(
                key="view_map",
                label="查看资产关系",
                description="确认它在账号、服务和系统之间的位置。",
                target="/map",
            ),
        ],
    )


def add_proposed_assignment(db: Session, asset: Asset, payload: ResourceCreate) -> None:
    if payload.owner_department_id:
        department = db.get(Department, payload.owner_department_id)
        if (
            department is None
            or department.archived_at is not None
            or (
                asset.legal_entity_id is not None
                and department.legal_entity_id != asset.legal_entity_id
            )
        ):
            raise HTTPException(status_code=422, detail="归属部门不属于当前公司")

    proposed_person_ids = {
        person_id
        for person_id in [payload.responsible_person_id, *payload.user_person_ids]
        if person_id is not None
    }
    if proposed_person_ids:
        people = list(
            db.scalars(
                select(Person).where(
                    Person.id.in_(proposed_person_ids), Person.archived_at.is_(None)
                )
            )
        )
        if len({person.id for person in people}) != len(proposed_person_ids) or (
            asset.legal_entity_id is not None
            and any(person.legal_entity_id != asset.legal_entity_id for person in people)
        ):
            raise HTTPException(status_code=422, detail="建议负责人和使用人员必须是当前公司成员")

    if payload.responsible_person_id:
        db.add(
            AssetResponsibility(
                asset_id=asset.id,
                person_id=payload.responsible_person_id,
                role_type="proposed_responsible",
                is_primary=True,
            )
        )
    for person_id in dict.fromkeys(payload.user_person_ids):
        if person_id != payload.responsible_person_id:
            db.add(
                AssetResponsibility(
                    asset_id=asset.id,
                    person_id=person_id,
                    role_type="proposed_user",
                    is_primary=False,
                )
            )


@router.post("/workspace/access-grants", status_code=status.HTTP_201_CREATED)
def create_access_grant(
    payload: AccessGrantCreate,
    db: Session = Depends(get_db),
    access: AccessContext = Depends(require_asset_write),
) -> dict[str, str]:
    if not payload.person_id and not payload.department_id:
        raise HTTPException(status_code=422, detail="人员和部门至少选择一个")
    target = db.get(Asset, payload.asset_id) if payload.asset_id else None
    if target is None and payload.account_id:
        account = db.get(Account, payload.account_id)
        if account is not None:
            target = db.get(Asset, account.asset_id)
    elif target is not None and payload.account_id:
        account = db.get(Account, payload.account_id)
        if account is None or account.asset_id != target.id:
            raise HTTPException(status_code=422, detail="账号与授权目标资产不一致")
    if target is None or target.archived_at is not None:
        raise HTTPException(status_code=422, detail="授权必须明确指向一个有效资产")
    require_asset_visible(db, access, target)
    if payload.person_id:
        person = db.get(Person, payload.person_id)
        if person is None or person.archived_at is not None:
            raise HTTPException(status_code=422, detail="所选人员不存在")
        if target.legal_entity_id and person.legal_entity_id != target.legal_entity_id:
            raise HTTPException(status_code=422, detail="人员与资产不属于同一公司")
    if payload.department_id:
        department = db.get(Department, payload.department_id)
        if department is None or department.archived_at is not None:
            raise HTTPException(status_code=422, detail="所选部门不存在")
        if target.legal_entity_id and department.legal_entity_id != target.legal_entity_id:
            raise HTTPException(status_code=422, detail="部门与资产不属于同一公司")
    grant_values = payload.model_dump()
    if grant_values.get("asset_id") is None:
        grant_values["asset_id"] = target.id
    item = AccessGrant(**grant_values, status="active")
    db.add(item)
    db.add(AuditLog(actor_user_id=access.user.id, action="asset.access_grant.create", object_type="asset", object_id=target.id, after_data={"person_id": str(payload.person_id) if payload.person_id else None, "department_id": str(payload.department_id) if payload.department_id else None}, request_id="workspace-access-grant"))
    db.commit()
    db.refresh(item)
    return {"id": str(item.id)}


@router.get("/asset-map", response_model=AssetMapRead)
def get_asset_map(
    include_people: bool = False,
    platform_id: UUID | None = None,
    db: Session = Depends(get_db),
    access: AccessContext = Depends(get_access_context),
) -> AssetMapRead:
    entities = list(db.scalars(select(LegalEntity).where(LegalEntity.archived_at.is_(None))))
    assets = list(
        db.scalars(
            select(Asset).where(Asset.archived_at.is_(None), asset_visibility_clause(access))
        )
    )
    visible_asset_ids = {item.id for item in assets}
    tenants = (
        list(
            db.scalars(
                select(PlatformTenant).where(
                    PlatformTenant.archived_at.is_(None),
                    PlatformTenant.asset_id.in_(visible_asset_ids),
                )
            )
        )
        if visible_asset_ids
        else []
    )
    service_instances = (
        list(
            db.scalars(
                select(ServiceInstance).where(
                    ServiceInstance.archived_at.is_(None),
                    ServiceInstance.asset_id.in_(visible_asset_ids),
                )
            )
        )
        if visible_asset_ids
        else []
    )
    account_identity_links = (
        list(
            db.scalars(
                select(PlatformAccountRegistrationIdentity).where(
                    PlatformAccountRegistrationIdentity.platform_tenant_id.in_(
                        {tenant.id for tenant in tenants}
                    ),
                    PlatformAccountRegistrationIdentity.status == "active",
                )
            )
        )
        if tenants
        else []
    )
    asset_platform_links = (
        list(
            db.scalars(
                select(AssetPlatformLink).where(
                    AssetPlatformLink.asset_id.in_(visible_asset_ids),
                    AssetPlatformLink.archived_at.is_(None),
                )
            )
        )
        if visible_asset_ids
        else []
    )
    visible_platform_ids = {item.platform_id for item in tenants} | {
        item.purchase_platform_id for item in service_instances if item.purchase_platform_id
    } | {item.platform_id for item in asset_platform_links}
    platforms = (
        list(
            db.scalars(
                select(Platform).where(
                    Platform.archived_at.is_(None), Platform.id.in_(visible_platform_ids)
                )
            )
        )
        if visible_platform_ids
        else []
    )
    visible_entity_ids = {item.legal_entity_id for item in assets if item.legal_entity_id}
    entities = [item for item in entities if item.id in visible_entity_ids]
    asset_types = {item.id: item for item in db.scalars(select(AssetType))}
    profiles = list(
        db.scalars(
            select(ResourceProfile).where(
                ResourceProfile.archived_at.is_(None),
                ResourceProfile.asset_id.in_(visible_asset_ids),
            )
        )
    )
    relations = list(
        db.scalars(
            select(AssetRelation).where(
                AssetRelation.archived_at.is_(None),
                AssetRelation.source_asset_id.in_(visible_asset_ids),
                AssetRelation.target_asset_id.in_(visible_asset_ids),
            )
        )
    )

    nodes: dict[str, MapNode] = {}
    edges: dict[str, MapEdge] = {}
    tenant_by_id = {item.id: item for item in tenants}
    tenant_asset_ids = {item.asset_id for item in tenants}
    service_instance_by_asset = {item.asset_id: item for item in service_instances}

    for entity in entities:
        key = f"entity:{entity.id}"
        nodes[key] = MapNode(id=key, label=entity.name, kind="company", status=entity.status)
    for platform in platforms:
        key = f"platform:{platform.id}"
        nodes[key] = MapNode(
            id=key,
            label=platform.name,
            kind="platform",
            status=platform.review_status,
            subtitle="平台",
        )
    for asset in assets:
        asset_type = asset_types.get(asset.asset_type_id)
        type_code = asset_type.code if asset_type else "generic"
        kind = (
            "platform_account" if asset.id in tenant_asset_ids else classify_asset_kind(type_code)
        )
        key = f"asset:{asset.id}"
        nodes[key] = MapNode(
            id=key,
            label=asset.name,
            kind=kind,
            status=asset.status,
            asset_id=asset.id,
            subtitle=asset_type.name if asset_type else None,
        )
    for tenant in tenants:
        edge_id = f"platform-tenant:{tenant.id}"
        edges[edge_id] = MapEdge(
            id=edge_id,
            source=f"platform:{tenant.platform_id}",
            target=f"asset:{tenant.asset_id}",
            relation="拥有账号",
        )
        entity_edge = f"entity-platform:{tenant.legal_entity_id}:{tenant.platform_id}"
        edges[entity_edge] = MapEdge(
            id=entity_edge,
            source=f"entity:{tenant.legal_entity_id}",
            target=f"platform:{tenant.platform_id}",
            relation="使用平台",
        )
    for link in account_identity_links:
        tenant = tenant_by_id.get(link.platform_tenant_id)
        if tenant is None:
            continue
        edges[f"identity-account:{link.id}"] = MapEdge(
            id=f"identity-account:{link.id}",
            source=f"asset:{link.registration_identity_asset_id}",
            target=f"asset:{tenant.asset_id}",
            relation="开通账号",
        )
    for link in asset_platform_links:
        edge_id = f"asset-platform-link:{link.id}"
        edges[edge_id] = MapEdge(
            id=edge_id,
            source=f"asset:{link.asset_id}",
            target=f"platform:{link.platform_id}",
            relation=relation_label(link.relation_type),
        )
    for asset in assets:
        service_instance = service_instance_by_asset.get(asset.id)
        if service_instance is None or service_instance.purchase_platform_id is None:
            continue
        purchase_platform_id = service_instance.purchase_platform_id
        if asset.legal_entity_id is not None:
            entity_platform_edge = f"entity-platform:{asset.legal_entity_id}:{purchase_platform_id}"
            edges.setdefault(
                entity_platform_edge,
                MapEdge(
                    id=entity_platform_edge,
                    source=f"entity:{asset.legal_entity_id}",
                    target=f"platform:{purchase_platform_id}",
                    relation="使用平台",
                ),
            )
        if service_instance.purchase_tenant_asset_id:
            edge_id = f"tenant-service:{service_instance.id}"
            edges[edge_id] = MapEdge(
                id=edge_id,
                source=f"asset:{service_instance.purchase_tenant_asset_id}",
                target=f"asset:{asset.id}",
                relation="购买 / 开通服务",
            )
            continue
        edges[f"platform-service:{service_instance.id}"] = MapEdge(
            id=f"platform-service:{service_instance.id}",
            source=f"platform:{purchase_platform_id}",
            target=f"asset:{asset.id}",
            relation="使用 / 开通服务",
        )
    for profile in profiles:
        if profile.managed_under_account_id in tenant_by_id:
            tenant = tenant_by_id[profile.managed_under_account_id]
            edge_id = f"resource-owner:{profile.id}"
            edges[edge_id] = MapEdge(
                id=edge_id,
                source=f"asset:{tenant.asset_id}",
                target=f"asset:{profile.asset_id}",
                relation="管理资源",
            )
    for relation in relations:
        source = f"asset:{relation.source_asset_id}"
        target = f"asset:{relation.target_asset_id}"
        label = relation_label(relation.relation_type)
        if any(
            edge.source == source and edge.target == target and edge.relation == label
            for edge in edges.values()
        ):
            continue
        edge_id = f"relation:{relation.id}"
        edges[edge_id] = MapEdge(
            id=edge_id,
            source=source,
            target=target,
            relation=label,
        )

    connected_assets = {
        endpoint.removeprefix("asset:")
        for edge in edges.values()
        for endpoint in (edge.source, edge.target)
        if endpoint.startswith("asset:")
    }
    for asset in assets:
        if str(asset.id) not in connected_assets and asset.legal_entity_id is not None:
            edge_id = f"entity-asset:{asset.id}"
            edges[edge_id] = MapEdge(
                id=edge_id,
                source=f"entity:{asset.legal_entity_id}",
                target=f"asset:{asset.id}",
                relation="归属",
            )

    if include_people:
        add_people_to_map(db, nodes, edges, visible_asset_ids)
    if platform_id is not None:
        nodes, edges = scope_asset_map_to_platform(nodes, edges, platform_id)
    referenced = {edge.source for edge in edges.values()} | {edge.target for edge in edges.values()}
    return AssetMapRead(
        nodes=[
            node
            for key, node in nodes.items()
            if key in referenced or node.kind == "company" or node.asset_id is not None
        ],
        edges=list(edges.values()),
    )


def scope_asset_map_to_platform(
    nodes: dict[str, MapNode], edges: dict[str, MapEdge], platform_id: UUID
) -> tuple[dict[str, MapNode], dict[str, MapEdge]]:
    """Return one platform's graph without using the company node as a cross-platform bridge."""
    platform_key = f"platform:{platform_id}"
    if platform_key not in nodes:
        raise HTTPException(status_code=404, detail="未找到或无权查看该平台")

    included = {platform_key}
    expandable = [platform_key]
    while expandable:
        current = expandable.pop()
        for edge in edges.values():
            if edge.source == current:
                other = edge.target
            elif edge.target == current:
                other = edge.source
            else:
                continue
            if other.startswith("entity:"):
                included.add(other)
                continue
            if other.startswith("platform:") and other != platform_key:
                continue
            if other not in included:
                included.add(other)
                expandable.append(other)

    return (
        {key: node for key, node in nodes.items() if key in included},
        {
            key: edge
            for key, edge in edges.items()
            if edge.source in included and edge.target in included
        },
    )


def classify_asset_kind(type_code: str) -> str:
    if type_code == "registration_identity":
        return "registration_identity"
    if type_code in {"internal_system", "automation_script", "ai_workflow"}:
        return "system"
    return "resource"


def relation_label(value: str) -> str:
    normalized = value.lower()
    return {
        "manages": "管理资源",
        "contains": "包含",
        "deployed_on": "部署于",
        "depends_on": "依赖",
        "uses": "使用",
        "registered_by": "注册于",
        "registered_on": "登记于平台",
        "provided_by": "由平台提供",
        "purchased_via": "购买于",
        "billed_to": "费用归集到",
        "stored_in": "存储于",
        "calls": "调用",
        "replaces": "替代",
    }.get(normalized, value)


def add_people_to_map(
    db: Session, nodes: dict[str, MapNode], edges: dict[str, MapEdge], visible_asset_ids: set[UUID]
) -> None:
    people = {
        item.id: item for item in db.scalars(select(Person).where(Person.archived_at.is_(None)))
    }
    departments = {item.id: item for item in db.scalars(select(Department))}
    responsibilities = list(
        db.scalars(
            select(AssetResponsibility).where(
                AssetResponsibility.archived_at.is_(None),
                AssetResponsibility.asset_id.in_(visible_asset_ids),
            )
        )
    )
    for item in responsibilities:
        if item.person_id and item.person_id in people:
            person = people[item.person_id]
            key = f"person:{person.id}"
            nodes[key] = MapNode(
                id=key,
                label=person.display_name,
                kind="person",
                status=person.employment_status,
                subtitle=departments.get(person.department_id).name
                if person.department_id in departments
                else None,
            )
            edge_id = f"responsibility:{item.id}"
            edges[edge_id] = MapEdge(
                id=edge_id,
                source=f"asset:{item.asset_id}",
                target=key,
                relation={"business_owner": "业务负责", "technical_owner": "技术负责"}.get(
                    item.role_type, item.role_type
                ),
            )


@router.post("/imports/analyze", response_model=FlexibleImportPreview)
async def analyze_flexible_import(
    file: Annotated[UploadFile | None, File()] = None,
    pasted_text: Annotated[str | None, Form()] = None,
) -> FlexibleImportPreview:
    if file is None and not pasted_text:
        raise HTTPException(status_code=422, detail="请选择文件或粘贴资料")
    if file is not None:
        content = await file.read()
        if len(content) > 20 * 1024 * 1024:
            raise HTTPException(status_code=413, detail="文件不能超过 20MB")
        source_name = file.filename or "未命名资料"
        source_kind, rows = parse_import_content(source_name, content)
        source_sha256 = hashlib.sha256(content).hexdigest()
    else:
        source_name = "粘贴资料"
        source_kind = "pasted_text"
        try:
            pasted_payload = json.loads(pasted_text or "")
        except json.JSONDecodeError:
            rows = [{"内容": line} for line in (pasted_text or "").splitlines() if line.strip()]
        else:
            source_kind = "json"
            rows = parse_legacy_json_payload(pasted_payload)
        source_sha256 = hashlib.sha256((pasted_text or "").encode("utf-8")).hexdigest()
    candidates = [suggest_candidate(index, row) for index, row in enumerate(rows, start=1)]
    return FlexibleImportPreview(
        source_name=source_name,
        source_kind=source_kind,
        source_sha256=source_sha256,
        recognized_counts=recognized_import_counts(candidates),
        candidates=candidates,
    )


def parse_import_content(file_name: str, content: bytes) -> tuple[str, list[dict]]:
    lower = file_name.lower()
    if lower.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        return "csv", list(csv.DictReader(io.StringIO(text)))
    if lower.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return "xlsx", []
        headers = [str(value or f"列{index + 1}").strip() for index, value in enumerate(values[0])]
        return "xlsx", [dict(zip(headers, row, strict=False)) for row in values[1:] if any(row)]
    if lower.endswith((".txt", ".md")):
        text = content.decode("utf-8-sig", errors="replace")
        return "text", [{"内容": line} for line in text.splitlines() if line.strip()]
    if lower.endswith(".json"):
        text = content.decode("utf-8-sig", errors="replace")
        try:
            payload = json.loads(text)
        except json.JSONDecodeError as error:
            raise HTTPException(status_code=422, detail="JSON 文件格式无效") from error
        return "json", parse_legacy_json_payload(payload)
    if lower.endswith((".html", ".htm")):
        text = content.decode("utf-8-sig", errors="replace")
        rows = parse_html_import_rows(text)
        if not rows:
            raise HTTPException(status_code=422, detail="HTML 中没有识别到可导入的业务表格")
        return "html", rows
    raise HTTPException(status_code=422, detail="当前支持 xlsx、csv、txt、md、html 或直接粘贴")


class BusinessTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.tables: list[list[list[str]]] = []
        self._table: list[list[str]] | None = None
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "table":
            self._table = []
        elif tag == "tr" and self._table is not None:
            self._row = []
        elif tag in {"th", "td"} and self._row is not None:
            self._cell = []

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag in {"th", "td"} and self._cell is not None and self._row is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None and self._table is not None:
            if any(self._row):
                self._table.append(self._row)
            self._row = None
        elif tag == "table" and self._table is not None:
            if self._table:
                self.tables.append(self._table)
            self._table = None


def parse_html_import_rows(text: str) -> list[dict]:
    parser = BusinessTableParser()
    parser.feed(text)
    relevant_headers = {
        "员工ID",
        "姓名",
        "昵称",
        "部门",
        "授权工具",
        "工具服务",
        "使用人",
        "平台",
        "账号",
        "域名",
        "服务器",
        "系统",
    }
    rows: list[dict] = []
    for table in parser.tables:
        headers = table[0]
        if not set(headers) & relevant_headers:
            continue
        for values in table[1:]:
            if any(values):
                rows.append(dict(zip(headers, values, strict=False)))

    # The historical FlipBelt AI distribution page renders data from a local
    # JavaScript object. Parse a fixed, data-only allowlist; never execute its JS.
    rows.extend(parse_legacy_flipbelt_arrays(text))
    return rows


LEGACY_FLIPBELT_ARRAYS = (
    ("assets", "legacy_asset", "ID", "employees"),
    ("employees", "access_grant", "员工ID", "transactions"),
    ("transactions", "expense_entry", "流水ID", "reimbursements"),
    ("reimbursements", "reimbursement_request", "报销ID", "payments"),
    ("payments", "payment_method_reference", "支付方式", "dingApps"),
    ("dingApps", "internal_system", "应用名称", "summary"),
    ("summary", "budget_snapshot", "month", "}"),
)

SENSITIVE_LEGACY_PAYMENT_FIELDS = {
    "cvv",
    "cvc",
    "securitycode",
    "security_code",
    "卡号",
    "银行卡号",
    "信用卡号",
    "cardnumber",
    "card_number",
}


def parse_legacy_flipbelt_arrays(text: str) -> list[dict]:
    rows: list[dict] = []
    for array_name, category, identifier_field, next_key in LEGACY_FLIPBELT_ARRAYS:
        rows.extend(
            parse_legacy_object_array(
                text,
                array_name=array_name,
                category=category,
                identifier_field=identifier_field,
                next_key=next_key,
            )
        )
    return rows


def parse_legacy_json_payload(payload: object) -> list[dict]:
    """Read a JSON export from FlipBelt localStorage without trusting executable code."""
    if isinstance(payload, list):
        return [item for item in payload if isinstance(item, dict)]
    if not isinstance(payload, dict):
        raise HTTPException(status_code=422, detail="JSON 导入内容必须是对象或对象数组")
    rows: list[dict] = []
    for array_name, category, identifier_field, _ in LEGACY_FLIPBELT_ARRAYS:
        collection = payload.get(array_name)
        if not isinstance(collection, list):
            continue
        for index, item in enumerate(collection, start=1):
            if not isinstance(item, dict):
                continue
            rows.append(
                annotate_legacy_import_row(
                    item,
                    array_name=array_name,
                    category=category,
                    identifier_field=identifier_field,
                    index=index,
                )
            )
    if rows:
        return rows
    return [payload]


def parse_legacy_employee_array(text: str) -> list[dict]:
    """Compatibility helper retained for callers of the original employee parser."""
    return parse_legacy_object_array(
        text,
        array_name="employees",
        category="access_grant",
        identifier_field="员工ID",
        next_key="transactions",
    )


def parse_legacy_object_array(
    text: str,
    *,
    array_name: str,
    category: str,
    identifier_field: str,
    next_key: str,
) -> list[dict]:
    rows: list[dict] = []
    end = re.escape(next_key) if next_key != "}" else r"}\s*;?"
    match = re.search(
        rf"\b{re.escape(array_name)}\s*:\s*\[(?P<body>.*?)\]\s*,?\s*(?:{end})\s*:",
        text,
        flags=re.DOTALL,
    )
    if next_key == "}":
        match = re.search(
            rf"\b{re.escape(array_name)}\s*:\s*\[(?P<body>.*?)\]\s*\}}\s*;?",
            text,
            flags=re.DOTALL,
        )
    if match is None:
        return []
    field_pattern = re.compile(
        r"(?P<key>[A-Za-z_\u4e00-\u9fff][\w\u4e00-\u9fff]*)\s*:\s*"
        r"(?P<value>\"(?:\\.|[^\"\\])*\"|-?\d+(?:\.\d+)?|true|false|null)"
    )
    for object_match in re.finditer(r"\{(?P<object>.*?)\}", match.group("body"), re.DOTALL):
        row: dict[str, object] = {}
        for field in field_pattern.finditer(object_match.group("object")):
            token = field.group("value")
            if token.startswith('"'):
                value: object = json.loads(token)
            elif token == "true":
                value = True
            elif token == "false":
                value = False
            elif token == "null":
                value = None
            else:
                value = float(token) if "." in token else int(token)
            key = field.group("key")
            row[key] = value
        if row:
            rows.append(
                annotate_legacy_import_row(
                    row,
                    array_name=array_name,
                    category=category,
                    identifier_field=identifier_field,
                    index=len(rows) + 1,
                )
            )
    return rows


def annotate_legacy_import_row(
    raw: dict,
    *,
    array_name: str,
    category: str,
    identifier_field: str,
    index: int,
) -> dict:
    row = dict(raw)
    sensitive_fields_removed = False
    if category == "payment_method_reference":
        for key in list(row):
            if normalize_legacy_field(str(key)) in SENSITIVE_LEGACY_PAYMENT_FIELDS:
                row.pop(key)
                sensitive_fields_removed = True
    identifier = str(row.get(identifier_field) or f"{array_name}:{index}").strip()
    row["__legacy_source_category"] = category
    row["__legacy_source_identifier"] = f"{array_name}:{identifier}"
    if sensitive_fields_removed:
        row["__legacy_sensitive_fields_removed"] = True
    return row


def normalize_legacy_field(value: str) -> str:
    return re.sub(r"[\s_-]", "", value).lower()


def suggest_candidate(row_number: int, raw: dict) -> ImportCandidate:
    source_category = raw.get("__legacy_source_category")
    source_identifier = raw.get("__legacy_source_identifier")
    sensitive_fields_removed = bool(raw.get("__legacy_sensitive_fields_removed"))
    normalized = {
        str(key): value for key, value in raw.items() if not str(key).startswith("__legacy_")
    }
    text = " ".join(str(value or "") for value in normalized.values()).strip()
    headers = " ".join(normalized).lower()
    lower = text.lower()
    object_type = legacy_object_type(source_category) or "resource"
    confidence = 0.55
    if source_category is not None:
        confidence = 0.98
    elif any(key in headers or key in lower for key in ["员工", "employee", "授权工具", "月费"]):
        object_type, confidence = "access_grant", 0.86
    elif "使用人" in normalized and "工具服务" in normalized:
        object_type, confidence = "access_grant", 0.84
    elif re.search(r"[^@\s]+@[^@\s]+\.[^@\s]+", text) or re.search(r"\b1\d{10}\b", text):
        object_type, confidence = "registration_identity", 0.76
    if (
        source_category is None
        and object_type != "access_grant"
        and any(name in lower for name in ["阿里云", "deepseek", "openai", "chatgpt", "腾讯云"])
    ):
        object_type, confidence = "platform_account", max(confidence, 0.82)
    if (
        source_category is None
        and object_type != "access_grant"
        and any(name in lower for name in ["ecs", "oss", "api", "服务器", "域名", "nas", "系统"])
    ):
        object_type, confidence = "resource", max(confidence, 0.84)
    if object_type == "access_grant":
        person = next(
            (
                str(normalized.get(key) or "").strip()
                for key in ["姓名", "昵称", "使用人"]
                if normalized.get(key)
            ),
            "待匹配人员",
        )
        tool = next(
            (
                str(normalized.get(key) or "").strip()
                for key in ["授权工具", "工具服务", "AI工具"]
                if normalized.get(key)
            ),
            "待匹配授权",
        )
        name = f"{person} · {tool}"
    elif source_category:
        name = legacy_candidate_name(source_category, normalized, row_number)
    else:
        name = next(
            (str(value).strip() for value in normalized.values() if value),
            f"第 {row_number} 条资料",
        )
    warnings = [] if text else ["该行没有可识别内容"]
    if source_category == "legacy_asset":
        warnings.append("需在审核时确认服务实例与账号引用的归并关系")
    if source_category in {"expense_entry", "reimbursement_request", "budget_snapshot"}:
        warnings.append("仅保留为待核对的财务历史，不计入当前费用统计")
    if source_category == "payment_method_reference":
        warnings.append("付款方式仅保留安全参考信息，不导入卡号或安全码")
    if sensitive_fields_removed:
        warnings.append("已剔除敏感付款字段")
    return ImportCandidate(
        row_number=row_number,
        source_category=str(source_category) if source_category else None,
        source_identifier=str(source_identifier) if source_identifier else None,
        suggested_object_type=object_type,
        suggested_name=name[:300],
        confidence=confidence,
        raw=sanitize_import_row(normalized),
        warnings=warnings,
    )


SENSITIVE_IMPORT_LABELS = (
    "密码", "password", "passwd", "cookie", "token", "secret", "api key",
    "apikey", "pin", "验证码", "安全码", "卡号", "cvv", "cvc",
)


def mask_import_email(value: str) -> str:
    match = re.fullmatch(r"([^@\s]+)@([^@\s]+)", value.strip())
    if not match:
        return value
    local, domain = match.groups()
    return f"{local[:2]}***@{domain}"


def mask_import_phone(value: str) -> str:
    normalized = value.strip()
    if re.fullmatch(r"1\d{10}", normalized):
        return f"{normalized[:3]}****{normalized[-4:]}"
    return value


def sanitize_import_value(key: str, value: object) -> object:
    if not isinstance(value, str):
        return value
    text = value.strip()
    label = key.casefold().replace(" ", "")
    if any(token in label for token in SENSITIVE_IMPORT_LABELS):
        return "[敏感内容已隐藏]"
    if re.fullmatch(r"[^@\s]+@[^@\s]+", text):
        return mask_import_email(text)
    if re.fullmatch(r"1\d{10}", text):
        return mask_import_phone(text)
    if key == "内容":
        if any(token in text.casefold() for token in SENSITIVE_IMPORT_LABELS):
            return "[敏感内容已隐藏]"
        if re.fullmatch(r"[^@\s]+@[^@\s]+", text):
            return mask_import_email(text)
        if re.fullmatch(r"1\d{10}", text):
            return mask_import_phone(text)
        if re.match(r"https?://", text, flags=re.IGNORECASE):
            return text.split("?", 1)[0].split("#", 1)[0]
        # Standalone password-like lines are common in text account sheets.
        if re.fullmatch(r"(?=.*[A-Za-z])(?=.*\d)[A-Za-z\d@!#$%^&*._-]{8,40}", text):
            return "[敏感内容已隐藏]"
    return value


def sanitize_import_row(row: dict[str, object]) -> dict[str, object]:
    return {str(key): sanitize_import_value(str(key), value) for key, value in row.items()}


def legacy_object_type(category: object) -> str | None:
    return {
        "legacy_asset": "service_instance",
        "access_grant": "access_grant",
        "expense_entry": "expense_entry",
        "reimbursement_request": "workflow_request",
        "payment_method_reference": "payment_method_reference",
        "internal_system": "internal_system",
        "budget_snapshot": "budget_snapshot",
    }.get(str(category))


def legacy_candidate_name(category: object, row: dict[str, object], row_number: int) -> str:
    values = {
        "legacy_asset": ("工具服务", "套餐用途", "ID"),
        "expense_entry": ("平台服务", "日期", "流水ID"),
        "reimbursement_request": ("申请人", "AI工具", "报销ID"),
        "payment_method_reference": ("支付方式", "用途"),
        "internal_system": ("应用名称", "平台", "APIID"),
        "budget_snapshot": ("月份", "month"),
    }.get(str(category), ())
    parts = [str(row.get(key)).strip() for key in values if row.get(key)]
    return " · ".join(parts)[:300] or f"第 {row_number} 条历史数据"


def recognized_import_counts(candidates: list[ImportCandidate]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for candidate in candidates:
        key = candidate.source_category or candidate.suggested_object_type
        counts[key] = counts.get(key, 0) + 1
    return counts


@router.post("/imports/stage", response_model=FlexibleImportStageResult)
def stage_flexible_import(
    payload: FlexibleImportStage,
    db: Session = Depends(get_db),
    _: AccessContext = Depends(require_asset_write),
) -> FlexibleImportStageResult:
    if not payload.candidates:
        raise HTTPException(status_code=422, detail="没有可保存的候选资料")
    batch = ImportBatch(
        file_name=payload.source_name,
        status="pending_review",
        total_count=len(payload.candidates),
        success_count=0,
        error_count=0,
        errors=(
            [
                {
                    "kind": "source_snapshot",
                    "sha256": payload.source_sha256,
                    "counts": payload.recognized_counts,
                }
            ]
            if payload.source_sha256
            else []
        ),
    )
    db.add(batch)
    db.flush()
    for candidate in payload.candidates:
        db.add(
            SourceImportRecord(
                import_batch_id=batch.id,
                source_kind=payload.source_kind,
                source_identifier=candidate.source_identifier or f"row:{candidate.row_number}",
                raw_payload=candidate.raw,
                suggested_object_type=candidate.suggested_object_type,
                suggested_name=candidate.suggested_name,
                confidence=Decimal(str(candidate.confidence)),
                mapping_status="pending_review",
                review_note="；".join(candidate.warnings) or None,
            )
        )
    db.flush()
    proposed_object_count, proposed_relation_count = rebuild_import_plan(db, batch)
    db.commit()
    return FlexibleImportStageResult(
        batch_id=batch.id,
        staged_count=len(payload.candidates),
        status=batch.status,
        proposed_object_count=proposed_object_count,
        proposed_relation_count=proposed_relation_count,
    )


@router.post("/imports/aliyun-colleague/commit")
def commit_aliyun_colleague_batch(
    db: Session = Depends(get_db),
    access: AccessContext = Depends(require_asset_write),
) -> dict[str, object]:
    """Commit the user-confirmed colleague batch using the account-first rules."""
    return commit_aliyun_colleague_import(db, actor_user_id=access.user.id)


@router.post("/imports/aliyun-colleague/reclassify-account-first")
def reclassify_aliyun_colleague_batch(
    db: Session = Depends(get_db),
    access: AccessContext = Depends(require_asset_write),
) -> dict[str, object]:
    return reclassify_aliyun_colleague_as_l4_accounts(db, actor_user_id=access.user.id)


@router.get(
    "/imports/{batch_id}/account-candidates",
    response_model=list[ImportAccountCandidateRead],
)
def list_import_account_candidates(
    batch_id: UUID,
    db: Session = Depends(get_db),
    _: AccessContext = Depends(require_asset_write),
) -> list[ImportAccountCandidateRead]:
    """Expose only redacted candidate fields kept in the import workbench."""
    batch = db.get(ImportBatch, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail="导入批次不存在")
    records = list(
        db.scalars(
            select(SourceImportRecord).where(
                SourceImportRecord.import_batch_id == batch_id,
                SourceImportRecord.archived_at.is_(None),
            )
        )
    )
    candidates: list[ImportAccountCandidateRead] = []
    for record in records:
        for item in (record.raw_payload or {}).get("account_candidates", []):
            if not isinstance(item, dict):
                continue
            identifier = item.get("login_identifier")
            platform_name = item.get("platform_name")
            if not isinstance(identifier, str) or not isinstance(platform_name, str):
                continue
            candidates.append(
                ImportAccountCandidateRead(
                    source_record_id=record.id,
                    platform_name=platform_name,
                    login_identifier=identifier,
                    source_file=record.source_identifier or "来源文件待确认",
                    evidence=[str(entry) for entry in item.get("evidence", [])],
                    confidence=float(item.get("confidence", record.confidence or 0)),
                    status=str(item.get("status", "待确认所属公司L4")),
                    pending_reason=str(item.get("pending_reason", "待人工确认")),
                    assigned_l4_asset_id=item.get("assigned_l4_asset_id"),
                    assigned_l4_name=item.get("assigned_l4_name"),
                )
            )
    return sorted(candidates, key=lambda item: (item.status != "待确认所属公司L4", item.platform_name, item.login_identifier))


@router.post("/imports/{batch_id}/rollback-aliyun-colleague")
def rollback_aliyun_colleague_batch(
    batch_id: UUID,
    db: Session = Depends(get_db),
    access: AccessContext = Depends(require_asset_write),
) -> dict[str, object]:
    try:
        return rollback_aliyun_colleague_import(db, batch_id=batch_id, actor_user_id=access.user.id)
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc


def import_plan_read(db: Session, batch: ImportBatch) -> ImportPlanRead:
    analysis = db.scalar(
        select(ImportAnalysis)
        .where(ImportAnalysis.import_batch_id == batch.id)
        .order_by(ImportAnalysis.created_at.desc())
    )
    if analysis is None:
        raise HTTPException(status_code=409, detail="该批次尚未生成资料规划")
    objects = list(
        db.scalars(
            select(ImportProposedObject)
            .where(ImportProposedObject.import_batch_id == batch.id)
            .order_by(ImportProposedObject.layer_code, ImportProposedObject.suggested_name)
        )
    )
    source_records = {
        item.id: item
        for item in db.scalars(
            select(SourceImportRecord).where(SourceImportRecord.import_batch_id == batch.id)
        )
    }
    relations = list(
        db.scalars(
            select(ImportProposedRelation)
            .where(ImportProposedRelation.import_batch_id == batch.id)
            .order_by(ImportProposedRelation.relation_type)
        )
    )
    return ImportPlanRead(
        batch_id=batch.id,
        analysis_mode=analysis.analysis_mode,
        catalog_fingerprint=analysis.catalog_fingerprint,
        summary=analysis.summary,
        objects=[
            ImportProposalObjectRead(
                id=item.id,
                proposal_key=item.proposal_key,
                source_record_id=item.source_import_record_id,
                source_file=batch.file_name,
                source_reference=(
                    source_records[item.source_import_record_id].source_identifier
                    if item.source_import_record_id in source_records
                    else None
                ),
                layer_code=item.layer_code,
                object_type=item.object_type,
                suggested_name=item.suggested_name,
                normalized_payload=item.normalized_payload,
                evidence=item.evidence,
                extraction_confidence=float(item.extraction_confidence),
                match_confidence=float(item.match_confidence),
                match_status=item.match_status,
                review_status=item.review_status,
                matched_asset_id=item.matched_asset_id,
                matched_platform_id=item.matched_platform_id,
                matched_person_id=item.matched_person_id,
                matched_department_id=item.matched_department_id,
                review_note=item.review_note,
            )
            for item in objects
        ],
        relations=[
            ImportProposalRelationRead(
                id=item.id,
                source_proposal_id=item.source_proposal_id,
                target_proposal_id=item.target_proposal_id,
                relation_type=item.relation_type,
                evidence=item.evidence,
                confidence=float(item.confidence),
                validation_status=item.validation_status,
                review_status=item.review_status,
                review_note=item.review_note,
            )
            for item in relations
        ],
    )


@router.get("/imports/batches", response_model=list[ImportBatchRead])
def list_import_batches(
    db: Session = Depends(get_db),
    _: AccessContext = Depends(require_asset_write),
) -> list[ImportBatchRead]:
    """List staged imports so a saved plan can be reopened from the workbench."""
    batches = list(
        db.scalars(select(ImportBatch).order_by(ImportBatch.created_at.desc()).limit(50))
    )
    result: list[ImportBatchRead] = []
    for batch in batches:
        analysis = db.scalar(
            select(ImportAnalysis)
            .where(ImportAnalysis.import_batch_id == batch.id)
            .order_by(ImportAnalysis.created_at.desc())
        )
        result.append(
            ImportBatchRead(
                id=batch.id,
                file_name=batch.file_name,
                status=batch.status,
                total_count=batch.total_count,
                proposed_object_count=len(
                    list(
                        db.scalars(
                            select(ImportProposedObject.id).where(
                                ImportProposedObject.import_batch_id == batch.id
                            )
                        )
                    )
                ),
                proposed_relation_count=len(
                    list(
                        db.scalars(
                            select(ImportProposedRelation.id).where(
                                ImportProposedRelation.import_batch_id == batch.id
                            )
                        )
                    )
                ),
                analysis_mode=analysis.analysis_mode if analysis else None,
                created_at=batch.created_at,
            )
        )
    return result


@router.get("/imports/{batch_id}/plan", response_model=ImportPlanRead)
def get_import_plan(
    batch_id: UUID,
    db: Session = Depends(get_db),
    _: AccessContext = Depends(require_asset_write),
) -> ImportPlanRead:
    batch = db.get(ImportBatch, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail="导入批次不存在")
    return import_plan_read(db, batch)


@router.post("/imports/{batch_id}/rebuild-plan", response_model=ImportPlanRead)
def rebuild_staged_import_plan(
    batch_id: UUID,
    db: Session = Depends(get_db),
    _: AccessContext = Depends(require_asset_write),
) -> ImportPlanRead:
    batch = db.get(ImportBatch, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail="导入批次不存在")
    if batch.status in {"pilot_imported", "completed"}:
        raise HTTPException(status_code=409, detail="已提交的批次不能重新规划")
    rebuild_import_plan(db, batch)
    db.commit()
    return import_plan_read(db, batch)


@router.post("/imports/{batch_id}/ai-analyze", response_model=ImportPlanRead)
def enhance_staged_import_with_ai(
    batch_id: UUID,
    db: Session = Depends(get_db),
    _: AccessContext = Depends(require_asset_write),
) -> ImportPlanRead:
    """Enrich source hints with DeepSeek, then rebuild the same review-only plan."""
    batch = db.get(ImportBatch, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail="导入批次不存在")
    if batch.status in {"pilot_imported", "completed"}:
        raise HTTPException(status_code=409, detail="已提交的批次不能重新分析")
    try:
        details = enhance_import_records_with_ai(db, batch)
    except AIImportError as exc:
        db.rollback()
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    rebuild_import_plan(
        db,
        batch,
        analysis_mode="ai_enhanced",
        analyzer_version=AI_ANALYZER_VERSION,
        provider="DeepSeek",
        model=get_settings().ai_import_model,
        analysis_details=details,
    )
    db.commit()
    return import_plan_read(db, batch)


@router.patch("/imports/{batch_id}/proposals/{proposal_id}", response_model=ImportPlanRead)
def review_import_proposal(
    batch_id: UUID,
    proposal_id: UUID,
    payload: ImportProposalObjectPatch,
    db: Session = Depends(get_db),
    access: AccessContext = Depends(require_asset_write),
) -> ImportPlanRead:
    batch = db.get(ImportBatch, batch_id)
    proposal = db.get(ImportProposedObject, proposal_id)
    if batch is None or proposal is None or proposal.import_batch_id != batch_id:
        raise HTTPException(status_code=404, detail="导入批次或候选对象不存在")
    if batch.status in {"pilot_imported", "completed"}:
        raise HTTPException(status_code=409, detail="已提交的批次不能修改规划")
    before = {"review_status": proposal.review_status, "review_note": proposal.review_note}
    for field, value in payload.model_dump(exclude_unset=True).items():
        setattr(proposal, field, value)
    db.add(
        AuditLog(
            actor_user_id=access.user.id,
            action="import.proposal.review",
            object_type="import_proposed_object",
            object_id=proposal.id,
            before_data=before,
            after_data={"review_status": proposal.review_status, "review_note": proposal.review_note, "matched_asset_id": str(proposal.matched_asset_id) if proposal.matched_asset_id else None, "matched_platform_id": str(proposal.matched_platform_id) if proposal.matched_platform_id else None},
            request_id="import-review",
        )
    )
    db.commit()
    return import_plan_read(db, batch)


@router.patch("/imports/{batch_id}/relations/{relation_id}", response_model=ImportPlanRead)
def review_import_relation(
    batch_id: UUID,
    relation_id: UUID,
    payload: ImportProposalRelationPatch,
    db: Session = Depends(get_db),
    access: AccessContext = Depends(require_asset_write),
) -> ImportPlanRead:
    batch = db.get(ImportBatch, batch_id)
    relation = db.get(ImportProposedRelation, relation_id)
    if batch is None or relation is None or relation.import_batch_id != batch_id:
        raise HTTPException(status_code=404, detail="导入批次或候选关系不存在")
    if batch.status in {"pilot_imported", "completed"}:
        raise HTTPException(status_code=409, detail="已提交的批次不能修改规划")
    if payload.review_status is not None and payload.review_status not in {"pending_review", "approved", "rejected"}:
        raise HTTPException(status_code=422, detail="关系审核状态无效")
    before = {"review_status": relation.review_status, "review_note": relation.review_note}
    for field, value in payload.model_dump(exclude_unset=True).items():
        setattr(relation, field, value)
    db.add(
        AuditLog(
            actor_user_id=access.user.id,
            action="import.proposal_relation.review",
            object_type="import_proposed_relation",
            object_id=relation.id,
            before_data=before,
            after_data={"review_status": relation.review_status, "review_note": relation.review_note},
            request_id="import-review",
        )
    )
    db.commit()
    return import_plan_read(db, batch)


@router.post("/imports/{batch_id}/dry-run", response_model=ImportDryRunResult)
def dry_run_staged_import(
    batch_id: UUID,
    db: Session = Depends(get_db),
    _: AccessContext = Depends(require_asset_write),
) -> ImportDryRunResult:
    batch = db.get(ImportBatch, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail="导入批次不存在")
    result = dry_run_import_plan(db, batch)
    return ImportDryRunResult(batch_id=batch.id, **result)


@router.post("/imports/{batch_id}/commit-reviewed", response_model=ImportCommitResult)
def commit_reviewed_import(
    batch_id: UUID,
    payload: ImportCommitRequest,
    db: Session = Depends(get_db),
    access: AccessContext = Depends(require_asset_write),
) -> ImportCommitResult:
    """Promote only explicitly approved L2/L3/L6 proposals; never infer L4/L5."""
    try:
        result = commit_reviewed_import_batch(
            db, batch_id=batch_id, actor_user_id=access.user.id, reason=payload.reason
        )
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    return ImportCommitResult(batch_id=batch_id, **{key: result[key] for key in ("status", "created", "reused", "skipped")})


@router.post("/imports/{batch_id}/commit-boss-pilot", response_model=BossPilotImportResult)
def commit_boss_pilot_import(
    batch_id: UUID,
    payload: BossPilotImportRequest,
    db: Session = Depends(get_db),
    access: AccessContext = Depends(require_asset_write),
) -> BossPilotImportResult:
    """Create a private, reviewable pilot owned by the nominated business owner.

    The caller remains the audit actor. The nominated boss is the business
    registrar and sole responsible person, so only that person and governance
    roles can see the newly created assets before employee grants are confirmed.
    """
    boss_person_id = payload.boss_person_id or access.person_id
    if boss_person_id is None:
        raise HTTPException(status_code=422, detail="请指定老板登记人，或由老板本人登录后提交")
    boss = db.get(Person, boss_person_id)
    legal_entity_id = payload.legal_entity_id or (boss.legal_entity_id if boss else None)
    if legal_entity_id is None:
        raise HTTPException(status_code=422, detail="无法确定本次导入的公司主体")
    require_legal_entity(db, legal_entity_id)
    if boss is None or boss.archived_at is not None or boss.legal_entity_id != legal_entity_id:
        raise HTTPException(status_code=422, detail="老板登记人必须属于所选公司主体")
    batch = db.get(ImportBatch, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail="导入批次不存在")
    if batch.status == "pilot_imported":
        raise HTTPException(status_code=409, detail="该批次已完成老板试点导入，不能重复导入")
    records = list(
        db.scalars(
            select(SourceImportRecord).where(
                SourceImportRecord.import_batch_id == batch_id,
                SourceImportRecord.archived_at.is_(None),
            )
        )
    )
    if not records:
        raise HTTPException(status_code=422, detail="批次中没有可导入的暂存记录")

    service_records = [item for item in records if item.suggested_object_type == "service_instance"]
    app_records = (
        [item for item in records if item.suggested_object_type == "internal_system"]
        if payload.include_internal_apps
        else []
    )
    access_records = [item for item in records if item.suggested_object_type == "access_grant"]
    finance_records = [
        item
        for item in records
        if item.suggested_object_type
        in {"expense_entry", "workflow_request", "payment_method_reference", "budget_snapshot"}
    ]
    if not service_records and not app_records:
        raise HTTPException(status_code=422, detail="批次中没有可创建的服务或内部应用")

    service_assets_by_legacy_id: dict[str, Asset] = {}
    service_assets_by_name: dict[str, list[Asset]] = {}
    service_assets_by_service_user: dict[tuple[str, str], list[Asset]] = {}
    identity_assets_by_fingerprint: dict[str, Asset] = {}
    for record in service_records:
        row = record.raw_payload
        service_name = legacy_value(row, "工具服务", "service", default=record.suggested_name)
        asset_type = require_asset_type(
            db,
            "api_service" if is_legacy_api_service(row) else "saas_subscription",
        )
        provider = get_or_create_legacy_provider(db, service_name)
        platform = get_or_create_legacy_platform(db, provider=provider, boss=boss)
        product = get_or_create_legacy_product(
            db,
            provider=provider,
            service_name=service_name,
            is_api=is_legacy_api_service(row),
        )
        asset = create_asset(
            db,
            name=legacy_service_asset_name(row, service_name),
            asset_type_id=asset_type.id,
            legal_entity_id=legal_entity_id,
            status_value="draft",
            description=legacy_asset_description(row),
            owner_department_id=None,
            ownership_scope="company",
            source_type="legacy_import",
            created_by_person_id=boss.id,
            review_status="pending_review",
        )
        db.add(
            ServiceInstance(
                asset_id=asset.id,
                service_product_id=product.id,
                purchase_platform_id=platform.id,
                subscription_name=legacy_value(row, "套餐用途", "plan"),
                currency="USD" if legacy_value(row, "月预算USD") else "CNY",
            )
        )
        make_boss_private_asset(db, access, asset, boss, record)
        account_reference = legacy_value(row, "账号引用", "account_ref")
        if is_explicit_registration_identity(account_reference):
            identity_asset = get_or_restore_legacy_registration_identity(
                db,
                access=access,
                batch=batch,
                legal_entity_id=legal_entity_id,
                boss=boss,
                identifier=account_reference,
                cache=identity_assets_by_fingerprint,
            )
            db.add(
                AssetRelation(
                    source_asset_id=asset.id,
                    target_asset_id=identity_asset.id,
                    relation_type="registered_by",
                    source_type="legacy_import",
                    note="由历史台账的账号引用建立；未据此推导公司平台账号。",
                )
            )
        record.canonical_asset_id = asset.id
        record.mapping_status = "pilot_imported"
        legacy_id = legacy_value(row, "ID", "id", default=record.source_identifier or "")
        if legacy_id:
            db.add(
                AssetIdentifier(
                    asset_id=asset.id,
                    # Source IDs are only unique inside their original import
                    # batch. Keeping that batch namespace preserves traceability
                    # without rejecting a later pilot import that has ID values
                    # such as API-001 as well.
                    namespace=f"flipbelt_ai_v4:{batch.id}:assets",
                    identifier_type="legacy_service_id",
                    identifier_value=legacy_id,
                    verification_status="pending",
                    source_import_record_id=record.id,
                    confidentiality="internal",
                )
            )
        service_assets_by_legacy_id[legacy_id] = asset
        service_assets_by_name.setdefault(service_name.casefold(), []).append(asset)
        source_user = legacy_value(row, "使用人", "user").casefold()
        if source_user:
            service_assets_by_service_user.setdefault(
                (service_name.casefold(), source_user), []
            ).append(asset)

    for record in app_records:
        row = record.raw_payload
        app_name = legacy_value(row, "应用名称", "appName", default=record.suggested_name)
        asset_type = require_asset_type(db, "internal_system")
        asset = create_asset(
            db,
            name=app_name,
            asset_type_id=asset_type.id,
            legal_entity_id=legal_entity_id,
            status_value="draft",
            description=legacy_value(row, "描述", "description"),
            owner_department_id=None,
            source_type="legacy_import",
            created_by_person_id=boss.id,
            review_status="pending_review",
        )
        db.add(InternalSystemProfile(asset_id=asset.id))
        make_boss_private_asset(db, access, asset, boss, record)
        record.canonical_asset_id = asset.id
        record.mapping_status = "pilot_imported"
        api_legacy_id = legacy_value(row, "APIID", "apiId")
        if api_asset := service_assets_by_legacy_id.get(api_legacy_id):
            db.add(
                AssetRelation(
                    source_asset_id=asset.id,
                    target_asset_id=api_asset.id,
                    relation_type="calls",
                    source_type="legacy_import",
                    note="由旧系统钉钉应用 APIID 映射，待老板确认",
                )
            )

    pending_access_grants = 0
    unresolved_access_grants = 0
    for record in access_records:
        row = record.raw_payload
        person_name = legacy_value(row, "姓名", "name")
        nickname = legacy_value(row, "昵称", "nickname")
        tool_name = legacy_value(row, "授权工具", "tool")
        person = resolve_legacy_person(
            db,
            legal_entity_id=legal_entity_id,
            person_name=person_name,
            nickname=nickname,
        )
        matching_assets: list[Asset] = []
        for source_user in (nickname, person_name):
            if source_user:
                matching_assets.extend(
                    service_assets_by_service_user.get(
                        (tool_name.casefold(), source_user.casefold()), []
                    )
                )
        # For legacy rows whose name and nickname are identical (for example
        # “归雾”), the same target is found twice.  De-duplicate before
        # deciding whether the service reference is unambiguous.
        matching_assets = list({asset.id: asset for asset in matching_assets}.values())
        if not matching_assets and tool_name:
            matching_assets = service_assets_by_name.get(tool_name.casefold(), [])
        matching_asset = matching_assets[0] if len(matching_assets) == 1 else None
        if person is None or matching_asset is None:
            unresolved_access_grants += 1
            record.mapping_status = "pending_review"
            record.review_note = "缺少可匹配人员或唯一服务实例，未创建正式授权"
            continue
        db.add(
            AccessGrant(
                asset_id=matching_asset.id,
                person_id=person.id,
                grant_type="legacy_ai_use",
                grant_role="member",
                status="pending_review",
                monthly_budget=legacy_decimal(row, "月费USD"),
                currency="USD" if legacy_value(row, "月费USD") else None,
                note="旧系统授权记录已暂存；确认人员与服务实例后才会生效",
                legacy_source_id=record.source_identifier,
                legacy_metadata=row,
            )
        )
        record.canonical_asset_id = matching_asset.id
        record.mapping_status = "pilot_imported_pending_review"
        pending_access_grants += 1

    for record in finance_records:
        record.mapping_status = "staged_finance_review"
        record.review_note = "财务对象尚未启用，保留原始记录，未写入费用或报销正式台账"

    batch.status = "pilot_imported"
    batch.success_count = (
        len(service_records)
        + len(app_records)
        + pending_access_grants
        + len(identity_assets_by_fingerprint)
    )
    batch.error_count = unresolved_access_grants
    batch.errors = [
        *batch.errors,
        {
            "kind": "pilot_import_summary",
            "boss_person_id": str(boss.id),
            "explicit_registration_identities": len(identity_assets_by_fingerprint),
            "staged_finance_records": len(finance_records),
            "unresolved_access_grants": unresolved_access_grants,
        },
    ]
    db.commit()
    return BossPilotImportResult(
        batch_id=batch.id,
        created_service_assets=len(service_records),
        created_internal_system_assets=len(app_records),
        pending_access_grants=pending_access_grants,
        unresolved_access_grants=unresolved_access_grants,
        staged_finance_records=len(finance_records),
        status=batch.status,
    )


@router.post(
    "/imports/{batch_id}/rollback-boss-pilot",
    response_model=BossPilotRollbackResult,
)
def rollback_boss_pilot_import(
    batch_id: UUID,
    db: Session = Depends(get_db),
    access: AccessContext = Depends(require_asset_write),
) -> BossPilotRollbackResult:
    """Precisely hide the canonical objects created by one boss-pilot batch.

    The source rows remain available for a corrected re-import.  Only records
    referenced by this batch are archived, so pre-existing inventory is never
    touched.
    """
    batch = db.get(ImportBatch, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail="导入批次不存在")
    if batch.status != "pilot_imported":
        raise HTTPException(status_code=409, detail="只有已试入库的批次可以回滚")

    records = list(
        db.scalars(
            select(SourceImportRecord).where(
                SourceImportRecord.import_batch_id == batch_id,
                SourceImportRecord.archived_at.is_(None),
            )
        )
    )
    asset_ids = {
        record.canonical_asset_id for record in records if record.canonical_asset_id is not None
    }
    derived_asset_ids = set(
        db.scalars(
            select(AssetIdentifier.asset_id).where(
                AssetIdentifier.namespace.like(f"flipbelt_ai_v4:{batch.id}:derived"),
                AssetIdentifier.archived_at.is_(None),
            )
        )
    )
    asset_ids.update(derived_asset_ids)
    assets = list(
        db.scalars(
            select(Asset).where(
                Asset.id.in_(asset_ids),
                Asset.source_type == "legacy_import",
                Asset.archived_at.is_(None),
            )
        )
    ) if asset_ids else []
    rollback_ids = {asset.id for asset in assets}
    now = datetime.now(UTC)

    grants = list(
        db.scalars(
            select(AccessGrant).where(
                AccessGrant.asset_id.in_(rollback_ids),
                AccessGrant.grant_type == "legacy_ai_use",
                AccessGrant.archived_at.is_(None),
            )
        )
    ) if rollback_ids else []
    relations = list(
        db.scalars(
            select(AssetRelation).where(
                (AssetRelation.source_asset_id.in_(rollback_ids))
                | (AssetRelation.target_asset_id.in_(rollback_ids)),
                AssetRelation.archived_at.is_(None),
            )
        )
    ) if rollback_ids else []

    for grant in grants:
        grant.archived_at = now
    for relation in relations:
        relation.archived_at = now
    responsibilities = (
        db.scalars(
            select(AssetResponsibility).where(
                AssetResponsibility.asset_id.in_(rollback_ids),
                AssetResponsibility.archived_at.is_(None),
            )
        )
        if rollback_ids
        else []
    )
    for item in responsibilities:
        item.archived_at = now
    identifiers = (
        db.scalars(
            select(AssetIdentifier).where(
                AssetIdentifier.asset_id.in_(rollback_ids),
                AssetIdentifier.archived_at.is_(None),
            )
        )
        if rollback_ids
        else []
    )
    for item in identifiers:
        item.archived_at = now
    service_instances = (
        db.scalars(
            select(ServiceInstance).where(
                ServiceInstance.asset_id.in_(rollback_ids),
                ServiceInstance.archived_at.is_(None),
            )
        )
        if rollback_ids
        else []
    )
    for item in service_instances:
        item.archived_at = now
    platform_tenants = (
        db.scalars(
            select(PlatformTenant).where(
                PlatformTenant.asset_id.in_(rollback_ids),
                PlatformTenant.archived_at.is_(None),
            )
        )
        if rollback_ids
        else []
    )
    for item in platform_tenants:
        item.archived_at = now
    registration_identities = (
        db.scalars(
            select(RegistrationIdentityProfile).where(
                RegistrationIdentityProfile.asset_id.in_(rollback_ids),
                RegistrationIdentityProfile.archived_at.is_(None),
            )
        )
        if rollback_ids
        else []
    )
    for item in registration_identities:
        item.archived_at = now
    for asset in assets:
        asset.archived_at = now
    for record in records:
        if record.canonical_asset_id in rollback_ids:
            record.canonical_asset_id = None
            record.mapping_status = "pending_review"
            record.review_note = "已按试入库批次回滚，可修正后重新提交"

    batch.status = "preview"
    batch.success_count = 0
    batch.error_count = 0
    batch.errors = [
        *batch.errors,
        {
            "kind": "pilot_import_rollback",
            "archived_assets": len(assets),
            "archived_access_grants": len(grants),
        },
    ]
    db.add(
        AuditLog(
            actor_user_id=access.user.id,
            action="legacy_import.boss_pilot.rollback",
            object_type="import_batch",
            object_id=batch.id,
            after_data={"archived_assets": len(assets), "archived_access_grants": len(grants)},
            request_id="legacy-boss-pilot-rollback",
        )
    )
    db.commit()
    return BossPilotRollbackResult(
        batch_id=batch.id,
        archived_assets=len(assets),
        archived_access_grants=len(grants),
        archived_relations=len(relations),
        status=batch.status,
    )


@router.post(
    "/imports/{batch_id}/normalize-service-hierarchy",
    response_model=LegacyHierarchyNormalizeResult,
)
def normalize_legacy_service_hierarchy(
    batch_id: UUID,
    db: Session = Depends(get_db),
    access: AccessContext = Depends(require_asset_write),
) -> LegacyHierarchyNormalizeResult:
    """Attach a legacy pilot's services to normalized platforms without inventing accounts."""
    batch = db.get(ImportBatch, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail="导入批次不存在")
    records = list(
        db.scalars(
            select(SourceImportRecord).where(
                SourceImportRecord.import_batch_id == batch_id,
                SourceImportRecord.suggested_object_type == "service_instance",
                SourceImportRecord.canonical_asset_id.is_not(None),
                SourceImportRecord.archived_at.is_(None),
            )
        )
    )
    linked = 0
    platform_ids: set[UUID] = set()
    for record in records:
        asset = db.get(Asset, record.canonical_asset_id)
        if asset is None:
            continue
        service_instance = db.scalar(
            select(ServiceInstance).where(
                ServiceInstance.asset_id == asset.id,
                ServiceInstance.archived_at.is_(None),
            )
        )
        if service_instance is None:
            continue
        service_name = legacy_value(record.raw_payload, "工具服务", "service", default=asset.name)
        provider = get_or_create_legacy_provider(db, service_name)
        product = db.get(ServiceProduct, service_instance.service_product_id)
        if product is not None and product.provider_id != provider.id:
            product.provider_id = provider.id
        boss = db.get(Person, asset.created_by_person_id) or db.get(Person, access.person_id)
        if boss is None:
            raise HTTPException(status_code=422, detail="无法确定历史服务的登记人")
        platform = get_or_create_legacy_platform(db, provider=provider, boss=boss)
        service_instance.purchase_platform_id = platform.id
        platform_ids.add(platform.id)
        linked += 1
    db.add(
        AuditLog(
            actor_user_id=access.user.id,
            action="legacy_import.hierarchy.normalize",
            object_type="import_batch",
            object_id=batch.id,
            after_data={"linked_service_instances": linked, "platform_count": len(platform_ids)},
            request_id="legacy-service-hierarchy-normalize",
        )
    )
    db.commit()
    return LegacyHierarchyNormalizeResult(
        batch_id=batch.id,
        linked_service_instances=linked,
        platform_count=len(platform_ids),
    )


def legacy_value(row: dict, *keys: str, default: str = "") -> str:
    for key in keys:
        value = row.get(key)
        if value is not None and str(value).strip():
            return str(value).strip()
    return default


def resolve_legacy_person(
    db: Session,
    *,
    legal_entity_id: UUID,
    person_name: str,
    nickname: str,
) -> Person | None:
    """Resolve old-system name/nickname without assuming old departments are current."""
    candidates = [value.strip().casefold() for value in (person_name, nickname) if value.strip()]
    if not candidates or any("待分配" in value for value in candidates):
        return None
    people = list(
        db.scalars(
            select(Person).where(
                Person.legal_entity_id == legal_entity_id,
                Person.archived_at.is_(None),
            )
        )
    )
    matches = []
    for person in people:
        aliases = {part.strip().casefold() for part in person.display_name.split("-")}
        aliases.add(person.display_name.strip().casefold())
        if any(value in aliases for value in candidates):
            matches.append(person)
    return matches[0] if len(matches) == 1 else None


def legacy_decimal(row: dict, *keys: str) -> Decimal | None:
    value = legacy_value(row, *keys)
    if not value:
        return None
    try:
        return Decimal(value)
    except ArithmeticError:
        return None


def is_legacy_api_service(row: dict) -> bool:
    # The purpose may say that a SaaS product has API access. That does not
    # make the purchased product itself an API service. Use the explicit row
    # type or the source ID instead of searching every free-text field.
    row_type = legacy_value(row, "类型", "type").casefold()
    legacy_id = legacy_value(row, "ID", "id").casefold()
    return "api" in row_type or bool(re.fullmatch(r"api[-_]\d+", legacy_id))


def get_or_create_legacy_provider(db: Session, service_name: str) -> Provider:
    provider_code, provider_name = legacy_platform_definition(service_name)
    provider = db.scalar(
        select(Provider).where(Provider.code == provider_code, Provider.archived_at.is_(None))
    )
    if provider is None:
        provider = Provider(
            code=provider_code,
            name=provider_name,
        )
        db.add(provider)
        db.flush()
    return provider


def legacy_platform_definition(service_name: str) -> tuple[str, str]:
    normalized = service_name.casefold().strip()
    for names, code, name in (
        (("chatgpt plus", "chatgpt team", "openai chatgpt"), "openai", "OpenAI"),
        (("claude pro", "anthropic claude"), "anthropic", "Anthropic"),
        (("google gemini",), "google", "Google"),
        (("yishangcloud",), "yishangcloud", "yishangcloud"),
        (("xai grok",), "xai", "xAI"),
        (("钉钉悟空",), "dingtalk", "钉钉"),
    ):
        if normalized in names:
            return code, name
    return f"legacy-{hashlib.sha256(service_name.encode('utf-8')).hexdigest()[:16]}", service_name


def get_or_create_legacy_platform(db: Session, *, provider: Provider, boss: Person) -> Platform:
    platform_code = f"{provider.code}-platform"
    platform = db.scalar(
        select(Platform).where(
            Platform.provider_id == provider.id,
            Platform.code == platform_code,
            Platform.archived_at.is_(None),
        )
    )
    if platform is None:
        platform = Platform(
            provider_id=provider.id,
            code=platform_code,
            name=provider.name,
            category="ai_platform",
            review_status="pending_review",
            description="由历史资料归并的平台；仅保留来源中明确出现的对象与关系。",
            submitted_by_person_id=boss.id,
        )
        db.add(platform)
        db.flush()
    return platform


def get_or_create_legacy_product(
    db: Session, *, provider: Provider, service_name: str, is_api: bool
) -> ServiceProduct:
    code = f"legacy-{hashlib.sha256(service_name.encode('utf-8')).hexdigest()[:16]}"
    product = db.scalar(
        select(ServiceProduct).where(
            ServiceProduct.provider_id == provider.id,
            ServiceProduct.code == code,
            ServiceProduct.archived_at.is_(None),
        )
    )
    if product is None:
        product = ServiceProduct(
            provider_id=provider.id,
            code=code,
            name=service_name,
            service_category="ai_api" if is_api else "ai_subscription",
            billing_mode="usage" if is_api else "subscription",
        )
        db.add(product)
        db.flush()
    return product


def legacy_service_asset_name(row: dict, service_name: str) -> str:
    plan = legacy_value(row, "套餐用途", "plan")
    legacy_id = legacy_value(row, "ID", "id")
    return " · ".join(part for part in [service_name, plan, legacy_id] if part)[:200]


def legacy_asset_description(row: dict) -> str:
    purpose = legacy_value(row, "套餐用途", "用途")
    note = legacy_value(row, "备注", "note")
    account_ref = legacy_value(row, "账号引用", "account_ref")
    details = [
        "由 FlipBelt AI 分发系统历史快照导入，待业务确认。",
        f"用途：{purpose}" if purpose else "",
        f"历史账号引用：{account_ref}" if account_ref else "",
        f"备注：{note}" if note else "",
    ]
    return "\n".join(part for part in details if part)


def is_explicit_registration_identity(value: str) -> bool:
    """Only source values that are identities may create layer 2 objects."""
    return infer_identity_type(value) in {"email", "phone", "wechat"}


def get_or_restore_legacy_registration_identity(
    db: Session,
    *,
    access: AccessContext,
    batch: ImportBatch,
    legal_entity_id: UUID,
    boss: Person,
    identifier: str,
    cache: dict[str, Asset],
) -> Asset:
    """Create one real identity per source value, without deriving layer 4."""
    value = identifier.strip()
    fingerprint = hashlib.sha256(value.lower().encode("utf-8")).hexdigest()
    if cached := cache.get(fingerprint):
        return cached

    profile = db.scalar(
        select(RegistrationIdentityProfile).where(
            RegistrationIdentityProfile.identifier_fingerprint == fingerprint
        )
    )
    identity_type = infer_identity_type(value)
    if profile is not None:
        asset = db.get(Asset, profile.asset_id)
        if asset is None or asset.legal_entity_id != legal_entity_id:
            raise HTTPException(status_code=422, detail="注册身份与当前公司主体不一致")
        if asset.archived_at is not None and asset.source_type == "legacy_import":
            asset.archived_at = None
            asset.name = value
            asset.description = "由历史台账账号引用登记；未据此推导公司平台账号。"
            asset.ownership_scope = "company"
            asset.created_by_person_id = boss.id
            asset.status = "draft"
            asset.review_status = "pending_review"
            profile.archived_at = None
            profile.identity_type = identity_type
            profile.identifier_value = value
            profile.identifier_masked = value
            profile.source_nature = "company_owned"
            profile.custodian_person_id = None
            profile.verification_status = "pending"
            for item in db.scalars(
                select(AssetIdentifier).where(AssetIdentifier.asset_id == asset.id)
            ):
                item.archived_at = None
        elif asset.archived_at is not None:
            raise HTTPException(status_code=422, detail="注册身份历史记录不可复用")
    else:
        asset_type = require_asset_type(db, "registration_identity")
        asset = create_asset(
            db,
            name=value,
            asset_type_id=asset_type.id,
            legal_entity_id=legal_entity_id,
            status_value="draft",
            description="由历史台账账号引用登记；未据此推导公司平台账号。",
            ownership_scope="company",
            source_type="legacy_import",
            created_by_person_id=boss.id,
            review_status="pending_review",
        )
        profile = RegistrationIdentityProfile(
            asset_id=asset.id,
            identity_type=identity_type,
            identifier_value=value,
            identifier_masked=value,
            identifier_fingerprint=fingerprint,
            source_nature="company_owned",
            verification_status="pending",
        )
        db.add(profile)

    marker = db.scalar(
        select(AssetIdentifier).where(
            AssetIdentifier.asset_id == asset.id,
            AssetIdentifier.namespace == f"flipbelt_ai_v4:{batch.id}:derived",
            AssetIdentifier.identifier_type == "legacy_registration_identity",
        )
    )
    if marker is None:
        db.add(
            AssetIdentifier(
                asset_id=asset.id,
                namespace=f"flipbelt_ai_v4:{batch.id}:derived",
                identifier_type="legacy_registration_identity",
                identifier_value=value,
                verification_status="pending",
                confidentiality="internal",
            )
        )
    else:
        marker.archived_at = None
        marker.identifier_value = value
    responsibility = db.scalar(
        select(AssetResponsibility).where(
            AssetResponsibility.asset_id == asset.id,
            AssetResponsibility.role_type == "responsible",
            AssetResponsibility.archived_at.is_(None),
        )
    )
    if responsibility is None:
        db.add(
            AssetResponsibility(
                asset_id=asset.id,
                person_id=boss.id,
                role_type="responsible",
                is_primary=True,
            )
        )
    db.add(
        AuditLog(
            actor_user_id=access.user.id,
            action="legacy_import.registration_identity",
            object_type="asset",
            object_id=asset.id,
            after_data={"source_batch_id": str(batch.id), "layer": 2},
            request_id="legacy-boss-pilot-import",
        )
    )
    cache[fingerprint] = asset
    return asset


def make_boss_private_asset(
    db: Session,
    access: AccessContext,
    asset: Asset,
    boss: Person,
    source_record: SourceImportRecord,
) -> None:
    db.add(
        AssetResponsibility(
            asset_id=asset.id,
            person_id=boss.id,
            role_type="responsible",
            is_primary=True,
        )
    )
    db.add(
        AuditLog(
            actor_user_id=access.user.id,
            action="legacy_import.boss_pilot",
            object_type="asset",
            object_id=asset.id,
            after_data={
                "created_by_person_id": str(boss.id),
                "owner_department_id": None,
                "review_status": "pending_review",
                "source_identifier": source_record.source_identifier,
            },
            request_id="legacy-boss-pilot-import",
        )
    )

import csv
import io
from typing import Annotated

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.access import AccessContext, require_asset_write
from app.db.session import get_db
from app.models import Asset, AssetType, Department, LegalEntity

router = APIRouter(tags=["data-transfer"])

HEADERS = [
    "资产名称",
    "资产类型编码",
    "公司编码",
    "部门编码",
    "状态",
    "重要程度",
    "说明",
]


class ImportRow(BaseModel):
    row_number: int
    name: str
    asset_type_id: str | None = None
    legal_entity_id: str | None = None
    owner_department_id: str | None = None
    status: str = "draft"
    criticality: str = "normal"
    description: str | None = None
    errors: list[str] = Field(default_factory=list)


class ImportCommit(BaseModel):
    file_name: str = "盘点导入.xlsx"
    rows: list[ImportRow]


def workbook_response(workbook: Workbook, file_name: str) -> StreamingResponse:
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{file_name}"'},
    )


@router.get("/imports/template")
def download_template() -> StreamingResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "资产盘点"
    sheet.append(HEADERS)
    sheet.append(
        [
            "示例：阿里云主账号",
            "platform_account",
            "COMPANY",
            "IT",
            "active",
            "critical",
            "请删除示例行",
        ]
    )
    sheet.freeze_panes = "A2"
    for column, width in zip("ABCDEFG", [28, 22, 16, 16, 14, 14, 40], strict=True):
        sheet.column_dimensions[column].width = width
    return workbook_response(workbook, "asset_inventory_template.xlsx")


@router.post("/imports/preview")
async def preview_import(
    file: Annotated[UploadFile, File()], db: Session = Depends(get_db)
) -> dict:
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="文件不能超过10MB")
    if file.filename and file.filename.lower().endswith(".csv"):
        text = content.decode("utf-8-sig")
        raw_rows = list(csv.DictReader(io.StringIO(text)))
    elif file.filename and file.filename.lower().endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        headers = [str(value or "").strip() for value in values[0]] if values else []
        raw_rows = [dict(zip(headers, row, strict=False)) for row in values[1:]]
    else:
        raise HTTPException(status_code=422, detail="仅支持.xlsx或.csv文件")

    types = {item.code: item for item in db.scalars(select(AssetType))}
    entities = {item.code: item for item in db.scalars(select(LegalEntity))}
    departments = {item.code: item for item in db.scalars(select(Department))}
    rows: list[dict] = []
    for index, raw in enumerate(raw_rows, start=2):
        name = str(raw.get("资产名称") or "").strip()
        type_code = str(raw.get("资产类型编码") or "").strip()
        entity_code = str(raw.get("公司编码") or "").strip()
        department_code = str(raw.get("部门编码") or "").strip()
        errors: list[str] = []
        if not name:
            errors.append("缺少资产名称")
        if type_code not in types:
            errors.append("资产类型编码不存在")
        if entity_code not in entities:
            errors.append("公司编码不存在")
        if department_code and department_code not in departments:
            errors.append("部门编码不存在")
        rows.append(
            ImportRow(
                row_number=index,
                name=name,
                asset_type_id=str(types[type_code].id) if type_code in types else None,
                legal_entity_id=str(entities[entity_code].id) if entity_code in entities else None,
                owner_department_id=(
                    str(departments[department_code].id) if department_code in departments else None
                ),
                status=str(raw.get("状态") or "draft"),
                criticality=str(raw.get("重要程度") or "normal"),
                description=str(raw.get("说明") or "").strip() or None,
                errors=errors,
            ).model_dump()
        )
    return {
        "file_name": file.filename or "import.xlsx",
        "rows": rows,
        "valid_count": sum(not row["errors"] for row in rows),
        "error_count": sum(bool(row["errors"]) for row in rows),
    }


@router.post("/imports/commit", deprecated=True)
def commit_import(
    payload: ImportCommit,
    _: AccessContext = Depends(require_asset_write),
) -> dict:
    raise HTTPException(
        status_code=409,
        detail="旧 Excel 直接提交已停用，请使用“资料接入工作台”完成六层规划、人工确认与提交预检。",
    )


@router.get("/exports/assets.xlsx")
def export_assets(db: Session = Depends(get_db)) -> StreamingResponse:
    assets = list(
        db.scalars(select(Asset).where(Asset.archived_at.is_(None)).order_by(Asset.asset_code))
    )
    types = {item.id: item.code for item in db.scalars(select(AssetType))}
    entities = {item.id: item.code for item in db.scalars(select(LegalEntity))}
    departments = {item.id: item.code for item in db.scalars(select(Department))}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "资产底库"
    sheet.append(["资产编号", *HEADERS, "最后更新"])
    for asset in assets:
        sheet.append(
            [
                asset.asset_code,
                asset.name,
                types.get(asset.asset_type_id, ""),
                entities.get(asset.legal_entity_id, ""),
                departments.get(asset.owner_department_id, ""),
                asset.status,
                asset.criticality,
                asset.description or "",
                asset.updated_at.isoformat(),
            ]
        )
    return workbook_response(workbook, "assets_export.xlsx")
